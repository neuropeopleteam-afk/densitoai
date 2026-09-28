#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
DensitoAI — пакетный инференс контроля качества DXA-снимков (ЛЦТ 2026).

Главные инженерные требования (ТЗ п. 2.7, рецензия Fable 5):
  * инференс НИКОГДА не падает целиком из-за одного плохого файла —
    каждый файл обрабатывается в собственном try/except, при сбое пишется
    строка-заглушка с processing_status="Failure";
  * выходной CSV имеет ТОЧНЫЕ имена колонок и ТОЧНЫЕ строки регионов /
    типов нарушений (см. config.yaml, секции regions/violations);
  * воспроизводимость: детерминированные модели, фиксированные сиды,
    никаких обращений во внешнюю сеть (веса бэкбона лежат в models/).

Архитектура (двухконтурная):
  Контур A — геометрия (geometry_features.py): угол оси, металл, укладка, ROI.
  Контур B — замороженный EfficientNet-B0 -> эмбеддинг -> PCA -> логрегрессия.
  Стэкинг — ранговое усреднение A и B относительно OOF-распределений обучения.
  Если .pkl-моделей нет — физические fallback-правила контура A (config.yaml).

Использование:
  python src/inference.py --input <папка_или_zip> --output <results.csv>
  python src/inference.py --input data/test --output out/results.csv --xlsx --debug-csv

Формат выходной строки (одна строка = один DICOM-файл):
  path_to_study, study_uid, image_uid, anatomical_region, quality_class,
  violation_type, quality_prob, processing_status, time_of_processing
"""
from __future__ import annotations

import warnings
warnings.filterwarnings("ignore")

import argparse
import csv
import hashlib
import json
import logging
import os
import pickle
import re
import shutil
import sys
import tempfile
import time
import traceback
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np

# --------------------------------------------------------------------------- #
# Пути проекта. Всё относительно корня репозитория, переопределяется env.
# --------------------------------------------------------------------------- #
SRC_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = Path(os.environ.get("DENSITO_ROOT", SRC_DIR.parent)).resolve()
MODELS_DIR = Path(os.environ.get("DENSITO_MODELS_DIR", PROJECT_ROOT / "models")).resolve()
CONFIG_PATH = Path(os.environ.get("DENSITO_CONFIG", PROJECT_ROOT / "config.yaml")).resolve()
sys.path.insert(0, str(SRC_DIR))

# Веса EfficientNet-B0 лежат в репозитории -> torchvision берёт их локально,
# без выхода в интернет (требование ТЗ п. 3.2). Должно быть ДО import torch.
_TORCH_HOME = MODELS_DIR / "torch_home"
if _TORCH_HOME.exists():
    os.environ.setdefault("TORCH_HOME", str(_TORCH_HOME))

import pydicom  # noqa: E402

from geometry_features import (  # noqa: E402
    PIXEL_SPACING_X_MM, PIXEL_SPACING_Y_MM,
    segment_bone, spine_axis_features, foreign_object_features,
    spine_positioning_features, hip_positioning_features,
)
from hip_features import hip_all_features  # noqa: E402
from calibration_utils import risk_level  # noqa: E402  (К3: правило уровня риска)
import preprocess  # noqa: E402  (инвариантная предобработка: маска тела, канонизация экспозиции)
import markup_clean  # noqa: E402  (очистка впечатанной разметки денситометра, docs/MARKUP_STRESS.md)

__version__ = "2.5.0"
LOG = logging.getLogger("densito.inference")
# pydicom шумит предупреждениями о нестандартных UID в анонимизированных файлах — не ошибка
logging.getLogger("pydicom").setLevel(logging.ERROR)

# --------------------------------------------------------------------------- #
# Конфиг (с жёстко зашитыми значениями по умолчанию на случай отсутствия yaml)
# --------------------------------------------------------------------------- #
HIP_POS_COLS = ["femur_solidity", "shaft_width_mm", "abs_shaft_angle_deg", "merge_height_mm",
                "medial_neck_extent_mm"]
HIP_ROI_COLS = ["scan_length_mm", "shaft_len_below_troch_mm"]
# Набор медиан для импутации NaN (ключи models/geometry_medians.json, версия 2.4.0). Не выводится
# из geometry_cols: список признаков бедра приведён к моделям, а поведение импутации не меняется.
IMPUTATION_MEDIAN_COLS = ["axis_angle_deg", "bone_area_ratio", "bone_width_ratio", "center_offset_ratio",
                          "edge_distance_ratio", "metal_metal_area_mm2", "metal_metal_max_intensity_gap",
                          "shaft_angle_deg"]

DEFAULT_CONFIG: Dict[str, Any] = {
    "version": "2.5.0",
    "output": {
        "columns": ["path_to_study", "study_uid", "image_uid", "anatomical_region",
                    "quality_class", "violation_type", "quality_prob",
                    "processing_status", "time_of_processing"],
        "status_success": "Success", "status_failure": "Failure",
        # строго < 0.5: строка Failure имеет quality_class 0, инвариант «класс 1 <=> prob >= 0.5»
        "violation_separator": ";", "fallback_quality_prob": 0.499999,
        "path_mode": "relative", "csv_encoding": "utf-8",
    },
    "regions": {"spine": "Поясничный отдел позвоночника",
                "hip": "Проксимальный отдел бедра",
                "default_when_unknown": "hip", "spine_min_cols": 290},
    "violations": {"sp_pos": "Некорректная укладка",
                   "sp_axis": "Не выравнена ось позвоночника",
                   "sp_art": "Присутствуют посторонние предметы",
                   "rh_pos": "Некорректная укладка", "rh_roi": "Некорректная область интереса",
                   "lh_pos": "Некорректная укладка", "lh_roi": "Некорректная область интереса"},
    "criteria_by_region": {"spine": ["sp_pos", "sp_axis", "sp_art"],
                           "right_hip": ["rh_pos", "rh_roi"],
                           "left_hip": ["lh_pos", "lh_roi"]},
    # Совпадает с feature_cols внутри models/model_*_geom.pkl и CRITERION_GEOMETRY_COLS в
    # src/train_stacked.py (бедро: hip_pos / hip_roi); сверка — tests/test_feature_contract.py.
    "geometry_cols": {"sp_pos": ["center_offset_ratio", "bone_width_ratio"],
                      "sp_axis": ["axis_angle_deg"],
                      "sp_art": ["metal_metal_band70_area_log", "metal_metal_band70_max_gap"],
                      "rh_pos": HIP_POS_COLS, "rh_roi": HIP_ROI_COLS,
                      "lh_pos": HIP_POS_COLS, "lh_roi": HIP_ROI_COLS},
    # H2 (2.4.0): признаки контура A из канонического эмбеддинга кадра (src/sppos_head.py); только позвоночник
    "features": {"sp_pos": ["synth_pos_logit"]},
    "stacking": {"weight_geom": 0.5, "weight_emb": 0.5, "weights_by_criterion": {},
                 "any_violation_aggregation": "max",
                 "any_blend_weight_model": 0.5, "consistent_quality_prob": True},
    "thresholds": {}, "fallback_threshold": 0.5,
    "fallback_rules": {  # значения синхронизированы с config.yaml (калибровка на трейне)
        "sp_axis": {"feature": "axis_angle_deg", "center": 3.6, "scale": 1.0, "direction": 1},
        "sp_art": {"feature": "metal_metal_area_mm2", "center": 670.0, "scale": 200.0, "direction": 1},
        "sp_pos": {"feature": "center_offset_ratio_abs", "center": 0.074, "scale": 0.02, "direction": 1},
        "rh_pos": {"feature": "shaft_angle_deg", "center": 16.3, "scale": 3.0, "direction": 1},
        "lh_pos": {"feature": "shaft_angle_deg", "center": 14.7, "scale": 3.0, "direction": 1},
        "rh_roi": {"feature": "edge_distance_mm", "center": 20.0, "scale": 3.0, "direction": -1},
        "lh_roi": {"feature": "edge_distance_mm", "center": 20.0, "scale": 3.0, "direction": -1},
        "decision_threshold": 0.5,
    },
    "validation": {"min_rows": 64, "max_rows": 4096, "min_cols": 64, "max_cols": 4096,
                   "min_bone_fraction": 0.01, "max_bone_fraction": 0.90},
    "pixel_spacing_mm": {"y": PIXEL_SPACING_Y_MM, "x": PIXEL_SPACING_X_MM},
}

DICOM_EXTENSIONS = {".dcm", ".dicom", ".dic", ".ima"}
# Предел вложенности архивов (zip в zip …): защита от архивных бомб.
# Архив глубже предела не распаковывается и попадает в выгрузку строкой Failure.
MAX_ZIP_DEPTH = int(os.environ.get("DENSITO_MAX_ZIP_DEPTH", "4"))
# Подсказки региона из имени файла (образец "Для теста.zip": CR000000_ПОП.dcm,
# CR000000_ППОБ.dcm, CR000001_ЛПОБ.dcm). На закрытом тесте суффиксов может не быть.
FILENAME_HINTS = (
    ("ППОБ", "right_hip"), ("ЛПОБ", "left_hip"), ("ПОП", "spine"),
    ("RIGHT", "right_hip"), ("LEFT", "left_hip"), ("SPINE", "spine"),
    ("_R_", "right_hip"), ("_L_", "left_hip"), ("LUMBAR", "spine"),
)


# Критерии ТЗ по сторонам бедра -> критерий единой модели бедра (ключ в stacking.weights_by_criterion,
# metrics_summary.json["hip"], models/model_hip_*.pkl). Веса стэкинга по стороне НЕ различаются.
HIP_CRIT_TO_MODEL = {"rh_pos": "hip_pos", "lh_pos": "hip_pos", "rh_roi": "hip_roi", "lh_roi": "hip_roi"}


def stacking_weights(cfg: Dict[str, Any], crit: str) -> Tuple[float, float]:
    """(w_geom, w_emb) для критерия. Приоритет: stacking.weights_by_criterion[<crit>] ->
    weights_by_criterion[<критерий единой модели бедра>] -> weight_geom / weight_emb (дефолт 0.5/0.5).
    Вес по критерию задаёт w_geom, w_emb = 1 - w_geom (вентиль К2, выбран nested CV,
    см. models/nested_gate_decisions.json и work/B/REPORT.md)."""
    st = cfg.get("stacking", {}) or {}
    by_crit = st.get("weights_by_criterion") or {}
    for key in (crit, HIP_CRIT_TO_MODEL.get(crit)):
        if key is not None and by_crit.get(key) is not None:
            wg = float(np.clip(float(by_crit[key]), 0.0, 1.0))
            return wg, 1.0 - wg
    return float(st.get("weight_geom", 0.5)), float(st.get("weight_emb", 0.5))


def preproc_variants(cfg: Dict[str, Any], crit: str) -> Dict[str, str]:
    """{'geom': вариант, 'emb': вариант} для критерия (К11).
    config.yaml: preprocessing.enabled + variant_by_criterion. Ключи rh_*/lh_* берут запись
    единой модели бедра (hip_pos / hip_roi). По умолчанию и при enabled: false — baseline (2.1.0)."""
    out = {"geom": "baseline", "emb": "baseline"}
    sec = cfg.get("preprocessing") or {}
    if not sec.get("enabled", False):
        return out
    by_crit = sec.get("variant_by_criterion") or {}
    for key in (crit, HIP_CRIT_TO_MODEL.get(crit)):
        entry = by_crit.get(key) if key is not None else None
        if isinstance(entry, dict):
            for k in ("geom", "emb"):
                if entry.get(k):
                    out[k] = str(entry[k])
            return out
    return out


def _deep_update(base: Dict[str, Any], upd: Dict[str, Any]) -> Dict[str, Any]:
    for k, v in (upd or {}).items():
        if isinstance(v, dict) and isinstance(base.get(k), dict):
            base[k] = _deep_update(base[k], v)
        else:
            base[k] = v
    return base


FAILURE_PROB_MAX = 0.499999


def failure_quality_prob(cfg_out: Dict[str, Any]) -> float:
    """quality_prob строки Failure: output.fallback_quality_prob, но строго < 0.5.

    У Failure quality_class 0, а инвариант выгрузки — «класс 1 <=> quality_prob >= 0.5».
    Значение конфига >= 0.5 (или нечисловое) заменяется на 0.499999 с предупреждением.
    """
    try:
        val = float(cfg_out.get("fallback_quality_prob", FAILURE_PROB_MAX))
    except (TypeError, ValueError):
        val = float("nan")
    if not (0.0 <= val < 0.5):
        LOG.warning("output.fallback_quality_prob=%r вне [0, 0.5) -> %s", cfg_out.get("fallback_quality_prob"),
                    FAILURE_PROB_MAX)
        return FAILURE_PROB_MAX
    return val


class ConfigError(RuntimeError):
    """config.yaml есть, но не читается (или нет PyYAML), либо его нет без явного разрешения."""


ALLOW_DEFAULT_CONFIG_ENV = "DENSITO_ALLOW_DEFAULT_CONFIG"


def load_config(path: Optional[Path] = None) -> Dict[str, Any]:
    """config.yaml поверх встроенных значений. Fail-closed: файл есть, но не читается (нет PyYAML,
    ошибка синтаксиса, не словарь, нет доступа) -> ConfigError с понятным текстом. Встроенные значения
    без файла — только при DENSITO_ALLOW_DEFAULT_CONFIG=1 (иначе ConfigError): молча работать на
    умолчаниях при сломанной поставке нельзя."""
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))  # deep copy
    path = Path(path) if path else CONFIG_PATH
    allow_defaults = os.environ.get(ALLOW_DEFAULT_CONFIG_ENV, "").strip() == "1"
    if not path.exists():
        if allow_defaults:
            LOG.warning("Config %s not found. Using built-in defaults (%s=1).", path, ALLOW_DEFAULT_CONFIG_ENV)
            return cfg
        raise ConfigError(f"файл конфигурации {path} не найден; укажите путь (--config или DENSITO_CONFIG) "
                          f"или, только для отладки, разрешите встроенные значения: {ALLOW_DEFAULT_CONFIG_ENV}=1")
    try:
        import yaml  # PyYAML
    except ImportError as e:
        raise ConfigError(f"файл конфигурации {path} есть, но PyYAML не установлен ({e}); "
                          f"установите PyYAML (requirements.txt)") from e
    try:
        with open(path, "r", encoding="utf-8") as f:
            user_cfg = yaml.safe_load(f)
    except Exception as e:  # noqa: BLE001
        raise ConfigError(f"файл конфигурации {path} не прочитан: {type(e).__name__}: {e}") from e
    if user_cfg is None:  # пустой файл — тоже сломанная поставка, а не повод молча взять умолчания
        if not allow_defaults:
            raise ConfigError(f"файл конфигурации {path} пуст; для отладки на встроенных значениях: "
                              f"{ALLOW_DEFAULT_CONFIG_ENV}=1")
        user_cfg = {}
    if not isinstance(user_cfg, dict):
        raise ConfigError(f"файл конфигурации {path}: ожидался словарь YAML, получено {type(user_cfg).__name__}")
    cfg = _deep_update(cfg, user_cfg)
    LOG.info("Config loaded: %s", path)
    return cfg


# --------------------------------------------------------------------------- #
# Утилиты
# --------------------------------------------------------------------------- #
def sigmoid(x: float) -> float:
    x = float(np.clip(x, -50, 50))
    return float(1.0 / (1.0 + np.exp(-x)))


def clip01(p: float) -> float:
    if p is None or not np.isfinite(p):
        return 0.5
    return float(min(1.0, max(0.0, float(p))))


def path_hash(path: Path) -> str:
    return hashlib.sha1(str(path).encode("utf-8", errors="ignore")).hexdigest()[:32]


_CONTENT_HASH_CACHE: Dict[str, str] = {}
_DIR_HASH_CACHE: Dict[str, str] = {}


def content_hash(path: Path) -> str:
    """sha256 байтов файла (первые 32 символа). Резервный идентификатор кадра не должен
    зависеть от имени файла: закрытый набор приходит с другими именами."""
    key = str(path)
    cached = _CONTENT_HASH_CACHE.get(key)
    if cached:
        return cached
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    out = h.hexdigest()[:32]
    _CONTENT_HASH_CACHE[key] = out
    return out


def dir_content_hash(directory: Path) -> str:
    """Идентификатор папки по содержимому её файлов (для исследования без StudyInstanceUID):
    sha1 отсортированных хэшей содержимого. Не зависит от имён файлов и папок.
    Для очень больших папок (> 500 файлов или > 1 ГБ) берётся отсортированный список размеров."""
    key = str(directory)
    cached = _DIR_HASH_CACHE.get(key)
    if cached:
        return cached
    try:
        files = sorted((p for p in directory.iterdir() if p.is_file()), key=lambda p: p.name)
        total = sum(p.stat().st_size for p in files)
        if len(files) > 500 or total > (1 << 30):
            parts = sorted(str(p.stat().st_size) for p in files)
        else:
            parts = sorted(content_hash(p) for p in files)
        out = hashlib.sha1("|".join(parts).encode("utf-8")).hexdigest()[:32]
    except Exception:  # noqa: BLE001
        out = path_hash(directory)
    _DIR_HASH_CACHE[key] = out
    return out


def config_hash(cfg: Dict[str, Any]) -> str:
    """Короткий sha256 конфигурации (та же формула, что в api_server._config_hash)."""
    return hashlib.sha256(json.dumps(cfg, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()[:12]


def is_dicom_candidate(path: Path) -> bool:
    """DICOM-кандидат: по расширению или по магическому числу 'DICM' (offset 128)."""
    if path.suffix.lower() in DICOM_EXTENSIONS:
        return True
    if path.suffix == "" or path.suffix.lower() not in {".csv", ".xlsx", ".txt", ".json",
                                                          ".png", ".jpg", ".jpeg", ".pdf",
                                                          ".zip", ".md", ".py", ".yaml"}:
        try:
            with open(path, "rb") as f:
                f.seek(128)
                return f.read(4) == b"DICM"
        except Exception:  # noqa: BLE001
            return False
    return False


def _fix_zip_name(zi: "zipfile.ZipInfo") -> str:
    """Имена в zip без флага UTF-8 Python декодирует как cp437; архивы с Windows с русскими
    именами обычно в cp866 (иногда cp1251). Восстанавливаем читаемое имя.

    Порядок (2.4.1, Б2):
      1) флаг UTF-8 (бит 11) выставлен — имя уже правильное;
      2) в архиве есть extra field 0x7075 (Info-ZIP Unicode Path): zipfile сам подставляет из него
         имя в `filename`, и оно отличается от `orig_filename` — берём как есть, повторная
         перекодировка превращала кириллицу в «???»;
      3) иначе перекодируем исходные байты имени (`orig_filename` в cp437) в cp866/cp1251."""
    if zi.flag_bits & 0x800:
        return zi.filename
    orig = zi.orig_filename.split("\x00", 1)[0]   # zipfile так же обрезает имя по нулевому байту
    if zi.filename != zi.orig_filename and zi.filename != orig.replace(os.sep, "/"):
        return zi.filename
    raw = orig.encode("cp437", errors="replace")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        pass

    def _score(txt: str) -> int:  # больше — правдоподобнее: кириллица/латиница хорошо, псевдографика плохо
        good = sum(1 for ch in txt if ch.isalnum() or ch in " _-./()")
        bad = sum(1 for ch in txt if 0x2500 <= ord(ch) <= 0x25FF or ord(ch) < 32)
        return good - 3 * bad

    best = zi.filename
    best_score = _score(best) - 1
    for enc in ("cp866", "cp1251"):
        try:
            cand = raw.decode(enc)
        except UnicodeDecodeError:
            continue
        sc = _score(cand)
        if sc > best_score:
            best, best_score = cand, sc
    return best


def unique_path(base: Path) -> Path:
    """Свободное имя рядом с `base`: `IM1.dcm` → `IM1 (2).dcm` → `IM1 (3).dcm` …
    Нужно, чтобы одинаковые имена входных файлов не перетирали друг друга: ТЗ п. 2.5
    требует строку на каждый входной файл."""
    if not base.exists():
        return base
    stem, suf = base.stem, base.suffix
    for i in range(2, 10000):
        cand = base.with_name(f"{stem} ({i}){suf}")
        if not cand.exists():
            return cand
    return base.with_name(f"{stem}_{uuid.uuid4().hex[:8]}{suf}")


# Пределы распаковки архива: защита от «zip-бомбы» (маленький архив, гигантское содержимое).
MAX_ZIP_MEMBERS = int(os.environ.get("DENSITO_MAX_ZIP_MEMBERS", "20000"))
MAX_ZIP_UNPACKED_MB = float(os.environ.get("DENSITO_MAX_ZIP_UNPACKED_MB", "4096"))
MAX_ZIP_RATIO = float(os.environ.get("DENSITO_MAX_ZIP_RATIO", "200"))


def safe_extract_zip(zip_path: Path, dest: Path) -> List[Tuple[Path, str]]:
    """Безопасная распаковка: нормализуем кодировку имён, отбрасываем абсолютные пути и
    `..` (zip-slip), служебные каталоги __MACOSX. Элементы с совпадающим полным именем
    (zip это допускает) получают различающиеся имена, а не перетирают друг друга.
    Возвращает [(файл на диске, путь внутри архива)]."""
    out: List[Tuple[Path, str]] = []
    dest = dest.resolve()
    cap = int(MAX_ZIP_UNPACKED_MB * 1024 * 1024)
    unpacked = 0
    with zipfile.ZipFile(zip_path) as zf:
        members = [zi for zi in zf.infolist() if not zi.is_dir()]
        if len(members) > MAX_ZIP_MEMBERS:
            raise ValueError(f"В архиве «{zip_path.name}» {len(members)} файлов — больше предела "
                             f"{MAX_ZIP_MEMBERS}. Разбейте архив на части.")
        declared = sum(int(zi.file_size or 0) for zi in members)
        if declared > cap:
            raise ValueError(f"Распакованный объём архива «{zip_path.name}» ({declared / 1048576:.0f} МБ) "
                             f"превышает предел {MAX_ZIP_UNPACKED_MB:.0f} МБ. Разбейте архив на части.")
        for zi in members:
            if zi.compress_size > 0 and zi.file_size > (1 << 20) \
                    and zi.file_size / zi.compress_size > MAX_ZIP_RATIO:
                raise ValueError(f"Элемент «{zi.filename}» архива «{zip_path.name}» распаковывается в "
                                 f"{zi.file_size / zi.compress_size:.0f}× больший объём (предел "
                                 f"{MAX_ZIP_RATIO:.0f}×) — архив отклонён как небезопасный.")
            name = _fix_zip_name(zi).replace("\\", "/")
            parts = [p for p in name.split("/") if p not in ("", ".", "..")]
            if not parts or parts[0] == "__MACOSX" or parts[-1] == ".DS_Store":
                continue
            rel = "/".join(parts)
            target = (dest / Path(*parts)).resolve()
            if dest not in target.parents:
                LOG.warning("Zip entry skipped (path escapes destination): %r", zi.filename)
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            if target.exists():
                target = unique_path(target)
                rel = "/".join(parts[:-1] + [target.name])
                LOG.warning("Zip entry %r duplicates an earlier name, stored as %r", zi.filename, rel)
            with zf.open(zi) as src, open(target, "wb") as dst:
                while True:
                    buf = src.read(1 << 20)
                    if not buf:
                        break
                    unpacked += len(buf)
                    if unpacked > cap:      # заявленный размер в заголовке может врать
                        raise ValueError(f"Распакованный объём архива «{zip_path.name}» превысил предел "
                                         f"{MAX_ZIP_UNPACKED_MB:.0f} МБ на файле «{rel}» — архив отклонён.")
                    dst.write(buf)
            out.append((target, rel))
    return out


def _is_own_series_zip(z: Path) -> bool:
    """Архив дополнительных серий, записанный DensitoAI (additional_series.zip с индексом series_index.csv)."""
    try:
        from series_zip import SERIES_ZIP_NAME, INDEX_NAME
        if z.name != SERIES_ZIP_NAME:
            return False
        with zipfile.ZipFile(z) as zf:
            with zf.open(INDEX_NAME) as f:
                return f.readline().decode("utf-8", "replace").startswith("zip_path,kind,study_uid,")
    except Exception:  # noqa: BLE001 — не наш архив или битый: обрабатывается как обычно
        return False


def _find_zips(root: Path) -> List[Path]:
    """Вложенные архивы в любом регистре расширения (.zip/.ZIP/.Zip), отсортированные."""
    return sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() == ".zip")


# Причины отказа распаковки по пути архива (верхнего уровня и вложенных): read_and_validate пишет
# их в строку Failure вместо общей фразы.
_BROKEN_ARCHIVES: Dict[Path, str] = {}


def discover_files(input_path: Path, tmp_holder: List[Path],
                   display: Optional[Dict[Path, str]] = None) -> Tuple[Path, List[Path]]:
    """Возвращает (корень, отсортированный список DICOM-кандидатов).
    Поддерживает папку, одиночный файл и zip-архив (распаковка во временную папку).
    В `display` (если передан) складывает отображаемый путь для файлов из архивов:
    `<путь архива относительно корня>/<путь внутри архива>` (для входного zip — только
    путь внутри архива), чтобы path_to_study не содержал временных каталогов."""
    input_path = Path(input_path)
    display = display if display is not None else {}
    if input_path.is_file() and input_path.suffix.lower() == ".zip":
        tmp_dir = Path(tempfile.mkdtemp(prefix="densito_in_"))
        tmp_holder.append(tmp_dir)
        try:
            extracted = safe_extract_zip(input_path, tmp_dir)
        except Exception as e:  # noqa: BLE001
            # Как у вложенного битого архива: пакет не обрывается, архив даёт одну строку Failure
            # (read_and_validate), причина — в debug CSV (error).
            LOG.warning("Input archive %s is broken: %s", input_path, e)
            _BROKEN_ARCHIVES[input_path.resolve()] = (
                f"Архив «{input_path.name}» повреждён, не является zip-файлом или превышает пределы "
                f"распаковки ({e}). Пересоздайте архив и загрузите снова.")
            display[input_path.resolve()] = input_path.name
            return input_path.parent, [input_path]
        for target, rel in extracted:
            display[target.resolve()] = rel
        LOG.info("Archive %s extracted to %s", input_path, tmp_dir)
        root = tmp_dir
    elif input_path.is_file():
        # Единое правило discovery: одиночный файл отбирается тем же is_dicom_candidate, что и файл
        # в папке (не-DICOM — например, .txt — строк не создаёт).
        return input_path.parent, ([input_path] if is_dicom_candidate(input_path) else [])
    else:
        root = input_path
    if not root.exists():
        raise FileNotFoundError(f"Input path does not exist: {input_path}")
    files = [p for p in sorted(root.rglob("*")) if p.is_file() and is_dicom_candidate(p)]
    # zip внутри папки и zip внутри zip — распаковываем на любой глубине до MAX_ZIP_DEPTH
    # (ограничение — защита от архивных бомб; архив глубже предела попадает в выгрузку
    # строкой Failure, а не исчезает молча)
    pending: List[Tuple[Path, str, int]] = []
    for z in _find_zips(root):
        if _is_own_series_zip(z):
            # 2.4.1: собственный архив дополнительных серий (выход прошлого запуска, если каталог результатов
            # лежит внутри входной папки) — не исследование; иначе повторный запуск получил бы лишние строки
            LOG.warning("Skipped own additional series archive (output of a previous run): %s", z)
            continue
        try:
            z_rel = str(z.resolve().relative_to(root.resolve()))
        except ValueError:
            z_rel = z.name
        pending.append((z, z_rel, 1))
    while pending:
        z, z_rel, depth = pending.pop(0)
        try:
            sub = Path(tempfile.mkdtemp(prefix="densito_in_"))
            tmp_holder.append(sub)
            extracted = safe_extract_zip(z, sub)
            for target, rel in extracted:
                display[target.resolve()] = f"{z_rel}/{rel}"
            files += [p for p in sorted(sub.rglob("*")) if p.is_file() and is_dicom_candidate(p)]
            inner = _find_zips(sub)
            for z2 in inner:
                try:
                    z2_rel = str(z2.resolve().relative_to(sub.resolve()))
                except ValueError:
                    z2_rel = z2.name
                if depth < MAX_ZIP_DEPTH:
                    pending.append((z2, f"{z_rel}/{z2_rel}", depth + 1))
                else:
                    LOG.warning("Archive nesting deeper than %d, not extracted: %s",
                                MAX_ZIP_DEPTH, f"{z_rel}/{z2_rel}")
                    display[z2.resolve()] = f"{z_rel}/{z2_rel}"
                    files.append(z2)
        except Exception as e:  # noqa: BLE001
            # битый архив не пропускаем молча: он попадёт в результаты строкой Failure
            LOG.warning("Nested archive %s is broken: %s", z, e)
            _BROKEN_ARCHIVES[z.resolve()] = (f"Архив «{z.name}» повреждён, не является zip-файлом или превышает "
                                             f"пределы распаковки ({e}). Пересоздайте архив и загрузите снова.")
            files.append(z)
    return root, files


# --------------------------------------------------------------------------- #
# Чтение и валидация DICOM
# --------------------------------------------------------------------------- #
@dataclass
class DicomInfo:
    ds: Any
    img_u8: np.ndarray
    rows: int
    cols: int
    study_uid: str
    image_uid: str
    pixel_spacing: Tuple[float, float]  # (y_mm, x_mm)
    warnings: List[str] = field(default_factory=list)
    tags: Dict[str, str] = field(default_factory=dict)
    exposure_gamma: float = 1.0  # γ канонизации экспозиции (1.0 — кадр уже в эталонной экспозиции)
    img_canonical: Optional[np.ndarray] = None  # кадр с канонизированной экспозицией (К11)


def _tag(ds, name: str, default: str = "") -> str:
    try:
        v = getattr(ds, name, None)
        if v is None:
            return default
        s = str(v).strip()
        return s if s else default
    except Exception:  # noqa: BLE001
        return default


def normalize_pixels_ex(ds) -> Tuple[np.ndarray, np.ndarray, float]:
    """(кадр baseline, кадр canonical, γ).

    baseline — ровно как в 2.1.0: MONOCHROME1 -> инверсия, Rescale, RGB -> gray,
    многокадровые -> первый кадр, перцентильное окно 1–99 %. Этот кадр идёт в оверлеи,
    SR, extras, OOD-gate и хэши — их цифры не меняются.
    canonical — тот же кадр с канонизацией экспозиции (src/preprocess.py); его берут только
    те критерии, для которых это выбрал nested (config.yaml: preprocessing.variant_by_criterion)."""
    arr = ds.pixel_array
    if arr is None or arr.size == 0:
        raise ValueError("empty pixel_array")
    arr = np.asarray(arr)
    if arr.ndim == 4:      # frames x H x W x C
        arr = arr[0]
    if arr.ndim == 3:
        if arr.shape[-1] in (3, 4) and arr.shape[0] > 4:   # RGB(A)
            arr = arr[..., :3].astype(np.float32).mean(axis=-1)
        else:                                               # frames x H x W
            arr = arr[0]
    if arr.ndim != 2:
        raise ValueError(f"unsupported pixel array shape {arr.shape}")
    arr = arr.astype(np.float32)
    if _tag(ds, "PhotometricInterpretation", "MONOCHROME2") == "MONOCHROME1":
        arr = arr.max() - arr
    try:
        slope = float(getattr(ds, "RescaleSlope", 1.0) or 1.0)
        intercept = float(getattr(ds, "RescaleIntercept", 0.0) or 0.0)
        arr = arr * slope + intercept
    except Exception:  # noqa: BLE001
        pass
    lo, hi = np.percentile(arr, [1, 99])
    if hi > lo:
        arr = np.clip((arr - lo) / (hi - lo) * 255.0, 0, 255)
    else:
        arr = np.zeros_like(arr)
    img_u8 = arr.astype(np.uint8)
    # Впечатанная разметка денситометра (контуры ROI, линии L1–L4, рамка шейки) — убирается до оценки
    # (src/markup_clean.py, docs/MARKUP_STRESS.md). Включается только при >= 4 тонких прямых отрезках:
    # на 499 кадрах заказчика не срабатывает ни разу, поэтому их кадры и числа поставки не меняются.
    img_u8, _mk = markup_clean.clean(img_u8)
    if _mk.get("markup_cleaned"):
        LOG.info("разметка денситометра на изображении убрана: отрезков %d, пикселей %d",
                 _mk["markup_segments"], _mk["markup_pixels"])
    # Канонизация экспозиции — тот же код, что у обучения (src/preprocess.py):
    # окно 1–99 % снимает только линейные сдвиги яркости, степенные — нет.
    img_canonical, gamma = preprocess.canonical_frame(img_u8)
    return img_u8, img_canonical, gamma


def normalize_pixels(ds) -> np.ndarray:
    """Кадр baseline (семантика 2.1.0) — используется аудитами в tools/ и хэшами."""
    return normalize_pixels_ex(ds)[0]


class UnsupportedInputError(ValueError):
    """Снимок вне области применения (src/region_support.py); строка уходит в Failure."""


REGION_SUPPORT_TAGS = ("Modality", "Manufacturer", "ManufacturerModelName", "BodyPartExamined",
                       "SeriesDescription", "ProtocolName", "StudyDescription")


def region_support_of(info: "DicomInfo") -> Tuple[bool, str]:
    """Проверка области применения по тегам DICOM и геометрии декодированного кадра (rows/cols
    пиксельного массива). Правила — src/region_support.check_tags; сомнение — в пользу обработки."""
    ok_tags, why_tags = region_support_tags_of(info)
    if not ok_tags:
        return ok_tags, why_tags
    return region_support_geometry_of(info)


def region_support_tags_of(info: "DicomInfo") -> Tuple[bool, str]:
    """Только явные противоречия в тегах: чужой аппарат, модальность, область или проекция."""
    from region_support import check_tags  # noqa: E402
    tags = {t: _tag(info.ds, t) for t in REGION_SUPPORT_TAGS}
    return check_tags({k: v for k, v in tags.items() if v}, cols=0, rows=0)


def region_support_geometry_of(info: "DicomInfo") -> Tuple[bool, str]:
    """Только геометрия кадра (наблюдаемый диапазон исследований заказчика)."""
    from region_support import check_tags  # noqa: E402
    return check_tags({}, cols=int(info.cols), rows=int(info.rows))


def read_and_validate(path: Path, cfg: Dict[str, Any]) -> DicomInfo:
    """Валидатор входа. Любая проблема -> исключение (обрабатывается выше как Failure)."""
    v = cfg["validation"]
    if path.suffix.lower() == ".zip":
        try:
            reason = _BROKEN_ARCHIVES.get(Path(path).resolve())
        except OSError:
            reason = None
        raise ValueError(reason or "Архив повреждён или не является zip-файлом; пересоздайте архив и загрузите снова")
    ds = pydicom.dcmread(str(path), force=True)
    if not hasattr(ds, "PixelData") and "PixelData" not in ds:
        raise ValueError("DICOM has no PixelData")
    # Сырой поток без преамбулы/file meta (некоторые архивы и PACS-экспорты): pydicom читает теги,
    # но для декодирования пикселей нужен TransferSyntaxUID -> восстанавливаем из фактической кодировки.
    if "TransferSyntaxUID" not in getattr(ds, "file_meta", {}):
        enc = getattr(ds, "original_encoding", (None, None))
        implicit_vr = enc[0] if enc[0] is not None else True
        little_endian = enc[1] if enc[1] is not None else True
        if not hasattr(ds, "file_meta") or ds.file_meta is None:
            ds.file_meta = pydicom.dataset.FileMetaDataset()
        if implicit_vr:
            ts = pydicom.uid.ImplicitVRLittleEndian
        elif little_endian:
            ts = pydicom.uid.ExplicitVRLittleEndian
        else:
            ts = pydicom.uid.ExplicitVRBigEndian
        ds.file_meta.TransferSyntaxUID = ts
    img, img_canonical, exposure_gamma = normalize_pixels_ex(ds)
    rows, cols = img.shape
    if not (v["min_rows"] <= rows <= v["max_rows"] and v["min_cols"] <= cols <= v["max_cols"]):
        raise ValueError(f"image size out of range: {rows}x{cols}")
    if int(img.max()) == int(img.min()):
        raise ValueError("constant (blank) image")

    warns: List[str] = []
    study_uid = _tag(ds, "StudyInstanceUID")
    image_uid = _tag(ds, "SOPInstanceUID")
    if not study_uid:
        # fallback: хэш СОДЕРЖИМОГО файлов родительской папки (обычно = папка исследования);
        # не зависит от имён — закрытый набор приходит с другими именами файлов и папок
        study_uid = "hash-" + dir_content_hash(path.parent)
        warns.append("no StudyInstanceUID -> hash of parent folder content")
    if not image_uid:
        image_uid = "hash-" + content_hash(path)
        warns.append("no SOPInstanceUID -> hash of file content")

    # PixelSpacing: читаем тег, иначе константа аппарата (с предупреждением)
    ps_y, ps_x = float(cfg["pixel_spacing_mm"]["y"]), float(cfg["pixel_spacing_mm"]["x"])
    for tag_name in ("PixelSpacing", "ImagerPixelSpacing"):
        val = getattr(ds, tag_name, None)
        if val is not None:
            try:
                ps_y, ps_x = float(val[0]), float(val[1])
                break
            except Exception:  # noqa: BLE001
                continue
    else:
        warns.append("no PixelSpacing tag -> device default 1.05x0.6 mm")

    tags = {k: _tag(ds, k) for k in ("Modality", "Manufacturer", "BodyPartExamined",
                                     "SeriesDescription", "ProtocolName", "Laterality",
                                     "ImageLaterality", "StudyDescription", "PhotometricInterpretation")}
    if tags["Modality"] and tags["Modality"] not in ("OT", "CR", "DX", "RG", "SC", ""):
        warns.append(f"unexpected Modality={tags['Modality']}")
    return DicomInfo(ds=ds, img_u8=img, rows=rows, cols=cols, study_uid=study_uid,
                     image_uid=image_uid, pixel_spacing=(ps_y, ps_x), warnings=warns, tags=tags,
                     exposure_gamma=exposure_gamma, img_canonical=img_canonical)


# --------------------------------------------------------------------------- #
# Определение региона
# --------------------------------------------------------------------------- #
def hip_side_by_density(img_u8: np.ndarray) -> str:
    """УСТАРЕВШИЙ резерв: сторона по распределению плотной кости (ошибается на ~24% снимков).
    Используется только если анатомический детектор недоступен."""
    h, w = img_u8.shape
    thr = np.percentile(img_u8, 60)
    left = (img_u8[:, : w // 2] > thr).mean()
    right = (img_u8[:, w // 2:] > thr).mean()
    return "right_hip" if right > left else "left_hip"


def detect_hip_side_region(img_u8: np.ndarray) -> Tuple[str, str]:
    """Сторона бедра -> internal region. Основной метод — анатомическое правило
    hip_features.detect_hip_side (таз всегда медиальнее диафиза); при любой ошибке —
    старая плотностная эвристика. Возвращает (region, source)."""
    try:
        from hip_features import detect_hip_side
        side = detect_hip_side(img_u8)
        return ("right_hip" if side == "right" else "left_hip"), "anatomy"
    except Exception as e:  # noqa: BLE001
        LOG.warning("detect_hip_side failed (%s) -> density heuristic", e)
        return hip_side_by_density(img_u8), "density"


def classify_region(info: DicomInfo, path: Path, cfg: Dict[str, Any]) -> Tuple[str, str]:
    """Возвращает (internal_region in {spine,right_hip,left_hip}, source).
    Приоритет: DICOM-теги -> ширина кадра (правило организаторов, стандартные 300/280/248 px)
    -> подсказка в имени файла (только для нестандартной ширины; организаторы подтвердили,
    что суффиксов в именах тестовых файлов не будет) -> ширина + анатомия стороны.
    Для нестандартной ширины без подсказки решение уточняет контентная модель (process_file)."""
    name_up = path.stem.upper()
    hint_region = next((region for hint, region in FILENAME_HINTS if hint.upper() in name_up), None)
    hint_src = next((f"filename:{hint}" for hint, region in FILENAME_HINTS if hint.upper() in name_up), "")
    std_cols = set(int(c) for c in cfg["regions"].get("standard_cols", [300, 280, 248]))

    text = " ".join([info.tags.get("BodyPartExamined", ""), info.tags.get("SeriesDescription", ""),
                     info.tags.get("ProtocolName", ""), info.tags.get("StudyDescription", "")]).upper()
    lat = (info.tags.get("Laterality") or info.tags.get("ImageLaterality") or "").upper()
    has_spine = any(k in text for k in ("SPINE", "LUMBAR", "LSPINE", "L-SPINE", "ПОЗВОНОЧ", "ПОП"))
    has_hip = any(k in text for k in ("HIP", "FEMUR", "БЕДР", "ПОБ"))
    # В описании маркеры обеих областей (например, «SPINE+HIP» в протоколе сеанса) -> теги не решают,
    # решает ширина кадра (300 px — позвоночник, 280/248 px — бедро), как без тегов.
    if has_spine and not has_hip:
        return "spine", "dicom_tags"
    if has_hip and not has_spine:
        if lat.startswith("R"):
            return "right_hip", "dicom_tags+laterality"
        if lat.startswith("L"):
            return "left_hip", "dicom_tags+laterality"
        region, src = detect_hip_side_region(info.img_u8)
        return region, f"dicom_tags+{src}"

    # Имя файла — только если ширина нестандартная (иначе решает содержимое)
    if hint_region is not None and info.cols not in std_cols:
        return hint_region, hint_src

    # Правило по ширине (подтверждено организаторами): 300 px спина, 280/248 бедро
    if info.cols >= int(cfg["regions"]["spine_min_cols"]):
        return "spine", "dims"
    if lat.startswith("R"):
        return "right_hip", "dims+laterality"
    if lat.startswith("L"):
        return "left_hip", "dims+laterality"
    region, src = detect_hip_side_region(info.img_u8)
    return region, f"dims+{src}"


def guess_region_without_pixels(path: Path, cfg: Dict[str, Any]) -> str:
    """Для строки-заглушки при сбое: заголовок DICOM (Columns) -> имя файла -> дефолт.
    Заголовок — первым, чтобы регион строки Failure не зависел от переименования файла (прогон-двойник, идея 1)."""
    try:
        ds = pydicom.dcmread(str(path), force=True, stop_before_pixels=True)
        cols = int(getattr(ds, "Columns", 0) or 0)
        if cols >= int(cfg["regions"]["spine_min_cols"]):
            return "spine"
        if cols > 0:
            return "left_hip"
    except Exception:  # noqa: BLE001
        pass
    name_up = path.stem.upper()
    for hint, region in FILENAME_HINTS:
        if hint.upper() in name_up:
            return region
    return "spine" if cfg["regions"]["default_when_unknown"] == "spine" else "left_hip"


def official_region_name(internal_region: str, cfg: Dict[str, Any]) -> str:
    return cfg["regions"]["spine"] if internal_region == "spine" else cfg["regions"]["hip"]


# --------------------------------------------------------------------------- #
# Признаки контура A
# --------------------------------------------------------------------------- #
def extract_geometry(info: DicomInfo, region: str, cfg: Dict[str, Any],
                     variant: str = "baseline") -> Dict[str, Any]:
    """Геометрические признаки (контур A) для заданного варианта предобработки (К11)."""
    with preprocess.variant(variant):
        return _extract_geometry_inner(info, region, cfg, variant)


def _extract_geometry_inner(info: DicomInfo, region: str, cfg: Dict[str, Any], variant: str) -> Dict[str, Any]:
    img = info.img_u8 if (variant == "baseline" or info.img_canonical is None) else info.img_canonical
    mask = segment_bone(img)
    feats: Dict[str, Any] = {"region": region}
    bone_fraction = float((mask > 0).mean())
    feats["bone_fraction"] = bone_fraction
    v = cfg["validation"]
    if not (v["min_bone_fraction"] <= bone_fraction <= v["max_bone_fraction"]):
        info.warnings.append(f"non-standard image: bone fraction {bone_fraction:.3f}")

    foreign = foreign_object_features(img, mask)
    feats.update({f"metal_{k}": v_ for k, v_ in foreign.items()})

    ps_y, ps_x = info.pixel_spacing
    if region == "spine":
        axis = spine_axis_features(img, mask)
        feats["axis_angle_deg"] = axis["axis_angle_deg"]
        feats["curvature"] = axis["curvature"]
        feats["valid_rows"] = axis["valid_rows"]
        feats.update(spine_positioning_features(img, mask))
        co = feats.get("center_offset_ratio")
        feats["center_offset_ratio_abs"] = abs(co) if co is not None else None
    else:
        feats.update(hip_positioning_features(img, mask))
        edr = feats.get("edge_distance_ratio")
        feats["edge_distance_mm"] = (edr * info.cols * ps_x) if edr is not None else None
        # Новые физически обоснованные признаки бедра (каноническая ориентация,
        # femur_solidity / shaft_width_mm / scan_length_mm / ...) — используются
        # моделями model_hip_pos_*.pkl / model_hip_roi_*.pkl. Считает свою маску
        # сегментации бедра внутри (segment_bone_hip), не переиспользует `mask`
        # от segment_bone (позвоночный сегментатор).
        try:
            feats.update(hip_all_features(img))
        except Exception as e:  # noqa: BLE001
            info.warnings.append(f"hip_all_features failed: {e}")
    return feats


# --------------------------------------------------------------------------- #
# Модели (pickle) и стэкинг
# --------------------------------------------------------------------------- #
class ModelBundle:
    """Обёртка над одним .pkl. Терпима к двум форматам:
       (а) объект с predict_proba (sklearn Pipeline);
       (б) dict {'scaler','pca'?,'clf','feature_cols'?,'medians'?,'oof_scores'?,'mirror_right'?}."""

    def __init__(self, obj: Any, name: str):
        self.name = name
        self.obj = obj
        self.meta: Dict[str, Any] = obj if isinstance(obj, dict) else {}

    def predict(self, X: np.ndarray) -> float:
        X = np.asarray(X, dtype=np.float64).reshape(1, -1)
        if isinstance(self.obj, dict):
            sc = self.obj.get("scaler")
            if sc is not None:
                X = sc.transform(X)
            pca = self.obj.get("pca")
            if pca is not None:
                X = pca.transform(X)
            clf = self.obj.get("clf") or self.obj.get("model")
            if clf is None:
                raise ValueError(f"{self.name}: dict without 'clf'")
            return float(clf.predict_proba(X)[0, 1])
        if hasattr(self.obj, "predict_proba"):
            return float(self.obj.predict_proba(X)[0, 1])
        if hasattr(self.obj, "decision_function"):
            return sigmoid(float(self.obj.decision_function(X)[0]))
        raise ValueError(f"{self.name}: unsupported model object {type(self.obj)}")


class ModelRegistry:
    """Загружает все доступные модели один раз. Отсутствие файла — warning, не ошибка."""

    def __init__(self, models_dir: Path, cfg: Dict[str, Any]):
        self.models_dir = Path(models_dir)
        self.cfg = cfg
        self.geom: Dict[Tuple[str, str], ModelBundle] = {}
        self.emb: Dict[Tuple[str, str], ModelBundle] = {}
        self.any_geom: Dict[str, ModelBundle] = {}
        self.any_emb: Dict[str, ModelBundle] = {}
        self.ref_geom: Dict[Tuple[str, str], np.ndarray] = {}
        self.ref_emb: Dict[Tuple[str, str], np.ndarray] = {}
        self.thresholds: Dict[str, float] = {}
        self.geometry_medians: Dict[str, float] = {}
        self.n_loaded = 0
        self._load()

    # ---- helpers
    def _load_pickle(self, path: Path) -> Optional[ModelBundle]:
        if not path.exists():
            return None
        try:
            with open(path, "rb") as f:
                obj = pickle.load(f)
            self.n_loaded += 1
            LOG.info("Model loaded: %s", path.name)
            return ModelBundle(obj, path.name)
        except Exception as e:  # noqa: BLE001
            LOG.warning("Model %s could not be loaded (%s) -> ignored", path.name, e)
            return None

    def _candidates(self, region: str, crit: str, kind: str) -> List[Path]:
        names = [f"model_{region}_{crit}_{kind}.pkl"]
        if region in ("right_hip", "left_hip"):
            base = crit.split("_", 1)[-1]  # pos / roi
            names += [f"model_hip_hip_{base}_{kind}.pkl", f"model_hip_{base}_{kind}.pkl"]
        return [self.models_dir / n for n in names]

    def _load(self):
        if not self.models_dir.exists():
            LOG.warning("Models dir %s does not exist -> geometric fallback rules only", self.models_dir)
        for region, crits in self.cfg["criteria_by_region"].items():
            for crit in crits:
                for kind, store in (("geom", self.geom), ("emb_pca", self.emb)):
                    for cand in self._candidates(region, crit, kind):
                        mb = self._load_pickle(cand)
                        if mb is not None:
                            store[(region, crit)] = mb
                            break
                    else:
                        LOG.warning("No %s model for %s/%s (looked for %s) -> fallback",
                                    kind, region, crit, self._candidates(region, crit, kind)[0].name)
                # референсные OOF-распределения для рангового усреднения
                self._load_reference(region, crit)
            for kind, store in (("geom", self.any_geom), ("emb_pca", self.any_emb)):
                mb = self._load_pickle(self.models_dir / f"model_{region}_any_{kind}.pkl")
                if mb is not None:
                    store[region] = mb
        self._load_thresholds()
        self._load_medians()
        self._load_calibration()

    def _load_calibration(self):
        """К3: models/calibration.pkl (MODEL_CONTRACT.md, раздел «calibration.pkl»): Platt по критерию и запас
        зоны «не уверен». Ключи хранятся как sp_*/hip_*; rh_*/lh_* берут hip_*. Отсутствие файла — не ошибка:
        p_cal = None, запаса нет (тогда «не уверен» только при fallback-правиле/отказе).
        config.yaml: uncertainty.margin_by_criterion[<crit>] (не null) перекрывает запас из pkl."""
        self.platt: Dict[str, Tuple[float, float]] = {}
        self.margins: Dict[str, Optional[float]] = {}
        calib: Dict[str, Any] = {}
        p = self.models_dir / "calibration.pkl"
        if p.exists():
            try:
                with open(p, "rb") as f:
                    calib = pickle.load(f)
                if not isinstance(calib, dict) or calib.get("kind") != "densito_calibration":
                    LOG.warning("calibration.pkl: неожиданный формат -> игнорирую")
                    calib = {}
                else:
                    LOG.info("Calibration loaded: calibration.pkl (%d criteria)", len(calib.get("criteria", {}) or {}))
            except Exception as e:  # noqa: BLE001
                LOG.warning("calibration.pkl unreadable: %s", e)
                calib = {}
        ucfg = self.cfg.get("uncertainty", {}) or {}
        if ucfg.get("enabled", True) is False:
            cfg_margins: Dict[str, Any] = {}
            calib_margins: Dict[str, Any] = {}
        else:
            cfg_margins = ucfg.get("margin_by_criterion", {}) or {}
            calib_margins = calib.get("margin_by_criterion", {}) or {}
        for region, crits in self.cfg["criteria_by_region"].items():
            for crit in crits:
                base = ("hip_" + crit.split("_", 1)[-1]) if region in ("right_hip", "left_hip") else crit
                rec = (calib.get("criteria", {}) or {}).get(base) or {}
                pl = rec.get("platt")
                if isinstance(pl, dict) and pl.get("a") is not None:
                    self.platt[crit] = (float(pl["a"]), float(pl["b"]))
                m = cfg_margins.get(crit, cfg_margins.get(base))
                if m is None:
                    m = calib_margins.get(base)
                self.margins[crit] = None if m is None else float(m)

    def _load_reference(self, region: str, crit: str):
        key = (region, crit)
        # приоритет: oof_scores внутри pickle
        for store, ref in ((self.geom, self.ref_geom), (self.emb, self.ref_emb)):
            mb = store.get(key)
            if mb is not None and isinstance(mb.meta.get("oof_scores"), (list, np.ndarray)):
                ref[key] = np.sort(np.asarray(mb.meta["oof_scores"], dtype=np.float64))
        cands = [self.models_dir / f"oof_stacked_{region}_{crit}.csv"]
        if region in ("right_hip", "left_hip"):
            base = crit.split("_", 1)[-1]
            cands += [self.models_dir / f"oof_stacked_hip_hip_{base}.csv", self.models_dir / f"oof_stacked_hip_{base}.csv"]
        path = next((c for c in cands if c.exists()), None)
        if path is None:
            return
        try:
            import pandas as pd
            df = pd.read_csv(path)
            if key not in self.ref_geom and "oof_geom" in df:
                self.ref_geom[key] = np.sort(df["oof_geom"].dropna().values.astype(np.float64))
            if key not in self.ref_emb and "oof_emb" in df:
                self.ref_emb[key] = np.sort(df["oof_emb"].dropna().values.astype(np.float64))
        except Exception as e:  # noqa: BLE001
            LOG.warning("OOF reference %s not loaded: %s", path.name, e)

    def _load_thresholds(self):
        summary: Dict[str, Any] = {}
        p = self.models_dir / "metrics_summary.json"
        if p.exists():
            try:
                with open(p, "r", encoding="utf-8") as f:
                    summary = json.load(f)
            except Exception as e:  # noqa: BLE001
                LOG.warning("metrics_summary.json unreadable: %s", e)
        for region, crits in self.cfg["criteria_by_region"].items():
            for crit in crits:
                t = self.cfg.get("thresholds", {}).get(crit)
                if t is None:  # порог, сохранённый внутри pickle (dict-формат)
                    for store in (self.geom, self.emb):
                        mb = store.get((region, crit))
                        if mb is not None and mb.meta.get("threshold") is not None:
                            t = mb.meta["threshold"]
                            break
                if t is None:
                    t = (summary.get(region, {}).get(crit, {}) or {}).get("threshold")
                if t is None and region in ("right_hip", "left_hip"):   # единая модель бедра
                    base = crit.split("_", 1)[-1]
                    hip = summary.get("hip", {}) or {}
                    for k in (f"hip_{base}", base, crit):
                        if isinstance(hip.get(k), dict) and hip[k].get("threshold") is not None:
                            t = hip[k]["threshold"]
                            break
                if t is None:
                    t = self.cfg.get("fallback_threshold", 0.5)
                    LOG.warning("No threshold for %s -> %.2f", crit, t)
                self.thresholds[crit] = float(t)

    def _load_medians(self):
        """Медианы геометрических признаков для импутации NaN (как при обучении).

        В образе лежит только `models/geometry_medians.json` — агрегаты без путей, UID и
        меток. `data/geometry_features.csv` (выгрузка обучающего набора) используется только
        на машине разработки, если JSON ещё не собран. Оба пути дают одни и те же числа
        (сверка: `python tools/make_geometry_medians.py --check`).
        """
        # Набор столбцов фиксирован (IMPUTATION_MEDIAN_COLS), а не берётся из geometry_cols: после
        # сверки geometry_cols с feature_cols моделей (A1, 24.09) импутация должна остаться прежней.
        # Контур A критерия импутирует по medians внутри pkl; здесь — только any-модели и запас.
        cols = list(IMPUTATION_MEDIAN_COLS)
        j = MODELS_DIR / "geometry_medians.json"
        if j.exists():
            try:
                payload = json.loads(j.read_text(encoding="utf-8"))
                med = payload.get("medians", payload) or {}
                for c in cols:
                    if c in med and med[c] is not None:
                        self.geometry_medians[c] = float(med[c])
            except Exception as e:  # noqa: BLE001
                LOG.warning("geometry medians json not loaded: %s", e)
        if len(self.geometry_medians) < len(cols):
            p = PROJECT_ROOT / "data" / "geometry_features.csv"
            if p.exists():
                try:
                    import pandas as pd
                    df = pd.read_csv(p, usecols=lambda c: c in cols or c == "region")
                    for c in cols:
                        if c not in self.geometry_medians and c in df:
                            self.geometry_medians[c] = float(np.nanmedian(df[c].values.astype(np.float64)))
                except Exception as e:  # noqa: BLE001
                    LOG.warning("geometry medians not loaded: %s", e)
        for c in cols:
            self.geometry_medians.setdefault(c, 0.0)

    # ---- inference helpers
    @staticmethod
    def percentile_rank(score: float, ref: Optional[np.ndarray]) -> Optional[float]:
        if ref is None or len(ref) == 0:
            return None
        # эквивалент pandas rank(pct=True) в пределе: доля референса <= score
        return float(np.searchsorted(ref, score, side="right") / len(ref))

    def geometry_vector(self, crit: str, feats: Dict[str, Any], mb: Optional[ModelBundle]) -> np.ndarray:
        cols = (mb.meta.get("feature_cols") if mb is not None else None) or self.cfg["geometry_cols"][crit]
        medians = (mb.meta.get("medians") if mb is not None else None) or {}
        if isinstance(medians, (list, np.ndarray)):
            medians = dict(zip(cols, medians))
        vec = []
        for c in cols:
            val = feats.get(c)
            if val is None or (isinstance(val, float) and not np.isfinite(val)):
                val = medians.get(c, self.geometry_medians.get(c, 0.0))
            vec.append(float(val))
        return np.asarray(vec, dtype=np.float64)

    def has_any_model(self) -> bool:
        return bool(self.geom or self.emb or self.any_geom or self.any_emb)


# --------------------------------------------------------------------------- #
# Контур B — эмбеддинги (ленивая инициализация, мягкий отказ)
# --------------------------------------------------------------------------- #
class EmbeddingExtractor:
    """Экстрактор эмбеддингов контура B. Поддерживает несколько источников весов одного
    EfficientNet-B0: ``imagenet`` (torchvision) и ``densito`` (models/backbone_densito.pth —
    наше GPU-предобучение на снимках кости; используется только моделями укладки, см.
    meta['emb_source'] в pkl и src/embeddings.py). Бэкбоны загружаются лениво, один раз."""

    DEFAULT_SOURCE = "imagenet"

    def __init__(self, enabled: bool = True):
        self.enabled = enabled
        self._backbones: Dict[str, Any] = {}
        self._failed: set = set()

    def _init(self, source: str = DEFAULT_SOURCE):
        if source in self._backbones or source in self._failed or not self.enabled:
            return
        try:
            import torch
            torch.manual_seed(0)
            torch.set_num_threads(max(1, min(8, os.cpu_count() or 1)))
            from embeddings import FrozenBackbone
            self._backbones[source] = FrozenBackbone(source)
            LOG.info("EfficientNet-B0 backbone '%s' ready (TORCH_HOME=%s)", source, os.environ.get("TORCH_HOME", "-"))
        except Exception as e:  # noqa: BLE001
            self._failed.add(source)
            if source == self.DEFAULT_SOURCE:
                LOG.warning("Embedding backbone unavailable (%s) -> contour B disabled", e)
            else:
                LOG.warning("Embedding backbone '%s' unavailable (%s) -> models with this source get no contour B", source, e)

    @property
    def _backbone(self):  # совместимость со старым кодом/тестами
        return self._backbones.get(self.DEFAULT_SOURCE)

    def extract(self, img_u8: np.ndarray, mirror: bool = False, source: str = DEFAULT_SOURCE) -> Optional[np.ndarray]:
        self._init(source)
        bb = self._backbones.get(source)
        if bb is None:
            return None
        try:
            img = img_u8[:, ::-1].copy() if mirror else img_u8
            return np.asarray(bb.extract(img), dtype=np.float32)
        except Exception as e:  # noqa: BLE001
            LOG.warning("Embedding extraction failed (%s): %s", source, e)
            return None

    def extract_many(self, img_u8: np.ndarray, sources, mirror: bool = False) -> Dict[str, np.ndarray]:
        """{source: embedding} для всех запрошенных источников (недоступные пропускаются)."""
        out: Dict[str, np.ndarray] = {}
        for src in sorted(set(sources)):
            e = self.extract(img_u8, mirror=mirror, source=src)
            if e is not None:
                out[src] = e
        return out


def _emb_for(mb, embs: Optional[Dict[str, np.ndarray]]) -> Optional[np.ndarray]:
    """Эмбеддинг нужного источника для модели (meta['emb_source'], по умолчанию imagenet)."""
    if not embs or mb is None:
        return None
    return embs.get(str(mb.meta.get("emb_source", EmbeddingExtractor.DEFAULT_SOURCE)))


# --------------------------------------------------------------------------- #
# Основной класс инференса
# --------------------------------------------------------------------------- #
class DensitoInference:
    def __init__(self, cfg: Optional[Dict[str, Any]] = None, models_dir: Optional[Path] = None,
                 use_embeddings: bool = True, visualize_dir: Optional[Path] = None,
                 sr_dir: Optional[Path] = None, roi_autocorrect_dir: Optional[Path] = None,
                 sr_study: bool = False, sr_study_dir: Optional[Path] = None, extras: bool = False,
                 sr_per_image: bool = False, sc_dir: Optional[Path] = None,
                 seg_dir: Optional[Path] = None):
        self.cfg = cfg or load_config()
        # [EXTRAS] экспериментальные флаги (work/D): белые линии, OOD-gate, эндопротез, когерентность исследования.
        # Не влияют на 9 колонок; пишутся в <output>_extras.csv и в self.last_extras_rows (API: details.extras).
        self.extras = bool(extras)
        self._extras_input: Optional[Dict[str, Any]] = None
        self.last_extras_rows: List[Dict[str, Any]] = []
        self.registry = ModelRegistry(Path(models_dir) if models_dir else MODELS_DIR, self.cfg)
        self.embedder = EmbeddingExtractor(enabled=use_embeddings)
        # H2 (2.4.0): голова «дефект укладки» для признака контура A sp_pos `synth_pos_logit`
        # (src/sppos_head.py). Грузится при старте, если признак нужен модели sp_pos или указан в
        # config.yaml → features; отсутствие файла — явная ошибка, а не молчаливый ноль.
        # При --no-embeddings эмбеддинга нет, признак взять негде -> подставляется медиана обучения
        # из pkl (в лог пишется предупреждение; режим без контура B и так деградированный).
        self.sppos_head = None
        self._sppos_head_warned = False
        if self._embedding_feature_cols():
            if use_embeddings:
                from sppos_head import SynthPosHead  # noqa: E402
                self.sppos_head = SynthPosHead(self.registry.models_dir / "head_densito_synth.pth")
                LOG.info("Synth pos head loaded: %s", self.sppos_head.path.name)
            else:
                LOG.warning("--no-embeddings: признак synth_pos_logit (sp_pos) не считается, "
                            "подставляется медиана обучения из pkl")
        # Резервный классификатор области (позвоночник/бедро) по эмбеддингам.
        # Используется ТОЛЬКО когда ширина снимка нестандартна (не 300/280/248 px),
        # т.е. правило организаторов по ширине неприменимо. OOF-точность 100 % (499 файлов).
        self.region_fallback = None
        rf = (Path(models_dir) if models_dir else MODELS_DIR) / "model_region_emb.pkl"
        if use_embeddings and rf.exists():
            try:
                with open(rf, "rb") as fh:
                    self.region_fallback = pickle.load(fh)
            except Exception as e:  # noqa: BLE001
                LOG.warning("Region fallback model not loaded: %s", e)
        if not self.registry.has_any_model():
            LOG.warning("No trained .pkl models found in %s -> using physical fallback rules "
                        "(config.yaml: fallback_rules). Retrain/save models to enable stacking.",
                        self.registry.models_dir)
        # --- Бонус-функции (ТЗ п.2.6): визуализация, DICOM SR, предложение коррекции области интереса.
        # Строго опциональны, по умолчанию выключены (None), не влияют на обязательный
        # CSV-выход даже при внутренней ошибке (каждая обёрнута в try/except в process_file).
        self.visualize_dir = Path(visualize_dir) if visualize_dir else None
        self.sr_dir = Path(sr_dir) if sr_dir else None
        self.roi_autocorrect_dir = Path(roi_autocorrect_dir) if roi_autocorrect_dir else None
        # --- Режим «один SR на исследование» (К10; методология НПКЦ ДиТ: нет SR / два SR — дефект).
        # Файлы пишутся после обработки всей партии в <каталог CSV>/sr/<study_uid>_SR.dcm
        # (или в sr_study_dir). Хэши оригинала (sha256 файла и пикселей) считаются всегда — они
        # попадают в debug-CSV и в SR; на официальный CSV режим не влияет.
        self.sr_study = bool(sr_study or sr_study_dir)
        # SR на снимок пишем только если его попросили явно ИЛИ если SR на исследование выключен:
        # иначе на одно исследование получается два набора SR (методика ЦДиТ ожидает один на исследование)
        self.sr_per_image = bool(sr_per_image)
        self.sc_dir = Path(sc_dir) if sc_dir else None
        if self.sc_dir:
            self.sc_dir.mkdir(parents=True, exist_ok=True)
        # --- Бонус «сегментация» (дополнительный функционал): экспорт масок структур как DICOM SEG + PNG + JSON
        # (src/segmentation_export.py). По умолчанию выключен; на 9 колонок и results.csv не влияет.
        self.seg_dir = Path(seg_dir) if seg_dir else None
        if self.seg_dir:
            self.seg_dir.mkdir(parents=True, exist_ok=True)
            LOG.info("Segmentation export enabled: seg_dir=%s", self.seg_dir)
        self.sr_study_dir = Path(sr_study_dir) if sr_study_dir else None
        self._study_headers: Dict[str, Dict[str, Any]] = {}
        self.last_study_sr: Dict[str, str] = {}
        # фактический файл на диске для каждой строки последнего run() (см. run, keep_temp)
        self.last_row_files: List[str] = []
        # {study_uid: {"spine": bool, "hip": bool, "note": str|None}} — полнота исследования по областям
        # (идея 4); считается по строкам результата после каждого run(), независимо от режима SR
        self.last_study_completeness: Dict[str, Dict[str, Any]] = {}
        if self.visualize_dir or self.sr_dir or self.roi_autocorrect_dir:
            for d in (self.visualize_dir, self.sr_dir, self.roi_autocorrect_dir):
                if d:
                    d.mkdir(parents=True, exist_ok=True)
            LOG.info("Bonus outputs enabled: visualize=%s sr=%s roi_autocorrect=%s",
                      self.visualize_dir, self.sr_dir, self.roi_autocorrect_dir)

    # ---- H2 (2.4.0): признаки контура A из эмбеддинга кадра
    def _embedding_feature_cols(self) -> Dict[str, List[str]]:
        """{критерий: [признаки из эмбеддинга]} — объединение config.yaml → features и feature_cols
        геометрических pkl (признак, который ждёт модель, считается обязательным)."""
        from sppos_head import FEATURE_NAME, REGIONS  # noqa: E402
        out: Dict[str, List[str]] = {}
        cfg_feats = self.cfg.get("features") or {}
        for region in REGIONS:
            for crit in self.cfg["criteria_by_region"].get(region, []):
                cols = set(cfg_feats.get(crit) or [])
                mb = self.registry.geom.get((region, crit))
                if mb is not None and FEATURE_NAME in (mb.meta.get("feature_cols") or []):
                    cols.add(FEATURE_NAME)
                if cols:
                    out[crit] = sorted(cols)
        return out

    def _add_embedding_features(self, region: str, info, crit_preproc: Dict[str, Dict[str, str]],
                                feats_by_variant: Dict[str, Dict[str, Any]],
                                embs_by_variant: Dict[str, Dict[str, np.ndarray]], mirror: bool) -> None:
        """Дописывает в признаки контура A критерия значения, посчитанные из эмбеддинга кадра.
        sp_pos/synth_pos_logit: логит головы на эмбеддинге densito канонического кадра — тот же эмбеддинг,
        что уже посчитан для контура B sp_pos (повторный проход бэкбона не нужен; если его нет —
        считается один раз здесь). Ошибка головы не глушится: строка уйдёт в Failure с причиной."""
        from sppos_head import FEATURE_NAME, EMB_SOURCE, EMB_VARIANT  # noqa: E402
        need = self._embedding_feature_cols()
        for crit in self.cfg["criteria_by_region"][region]:
            cols = need.get(crit)
            if not cols or FEATURE_NAME not in cols:
                continue
            geom_variant = crit_preproc[crit]["geom"]
            target = feats_by_variant.setdefault(geom_variant, {})
            if self.sppos_head is None:
                if not self._sppos_head_warned:
                    LOG.warning("%s: признак %s не посчитан (контур B выключен) -> медиана обучения", crit, FEATURE_NAME)
                    self._sppos_head_warned = True
                target[FEATURE_NAME] = None
                continue
            emb = (embs_by_variant.get(EMB_VARIANT) or {}).get(EMB_SOURCE)
            if emb is None:
                frame = info.img_canonical if (EMB_VARIANT == "canonical" and info.img_canonical is not None) else info.img_u8
                emb = self.embedder.extract(frame, mirror=mirror, source=EMB_SOURCE)
                if emb is None:
                    raise RuntimeError(f"{crit}: не удалось получить эмбеддинг '{EMB_SOURCE}'/'{EMB_VARIANT}' "
                                       f"для признака {FEATURE_NAME}")
                embs_by_variant.setdefault(EMB_VARIANT, {})[EMB_SOURCE] = emb
            target[FEATURE_NAME] = self.sppos_head.logit(emb)

    # ---- скоринг одного критерия
    def score_criterion(self, region: str, crit: str, feats: Dict[str, Any],
                        embs: Optional[Dict[str, np.ndarray]]) -> Dict[str, Any]:
        reg = self.registry
        key = (region, crit)
        mb_g, mb_e = reg.geom.get(key), reg.emb.get(key)
        out: Dict[str, Any] = {"p_geom": None, "p_emb": None, "rank_geom": None,
                               "rank_emb": None, "score": None, "method": None}

        if mb_g is not None:
            try:
                out["p_geom"] = clip01(mb_g.predict(reg.geometry_vector(crit, feats, mb_g)))
                out["rank_geom"] = reg.percentile_rank(out["p_geom"], reg.ref_geom.get(key))
            except Exception as e:  # noqa: BLE001
                LOG.warning("%s/%s geom model failed: %s", region, crit, e)
        emb = _emb_for(mb_e, embs)
        if mb_e is not None and emb is not None:
            try:
                out["p_emb"] = clip01(mb_e.predict(emb))
                out["rank_emb"] = reg.percentile_rank(out["p_emb"], reg.ref_emb.get(key))
            except Exception as e:  # noqa: BLE001
                LOG.warning("%s/%s emb model failed: %s", region, crit, e)

        wg, we = stacking_weights(self.cfg, crit)
        out["w_geom"] = wg
        rg = out["rank_geom"] if out["rank_geom"] is not None else out["p_geom"]
        re_ = out["rank_emb"] if out["rank_emb"] is not None else out["p_emb"]
        if rg is not None and re_ is not None:
            # вес по критерию (вентиль): при wg=1 скор = rank_geom, при wg=0 — rank_emb
            out["score"], out["method"] = wg * rg + we * re_, "stacked_rank_avg"
        elif rg is not None:
            out["score"], out["method"] = rg, "geom_only"
        elif re_ is not None:
            out["score"], out["method"] = re_, "emb_only"
        else:
            out["score"], out["method"] = self.fallback_rule(crit, feats), "fallback_rule"

        thr = (reg.thresholds.get(crit, 0.5) if out["method"] != "fallback_rule"
               else float(self.cfg["fallback_rules"].get("decision_threshold", 0.5)))
        out["threshold"] = thr
        out["score"] = clip01(out["score"])
        out["flag"] = int(out["score"] >= thr)
        # К3 (только debug/API, на 9 колонок не влияет): запас до порога, зона «не уверен», Platt-вероятность.
        # «не уверен» <=> |score - threshold| <= margin (включительно); при fallback-правиле — всегда «не уверен»,
        # при отсутствии запаса (нет calibration.pkl) — только при fallback.
        margin = getattr(reg, "margins", {}).get(crit)
        out["margin"] = round(abs(out["score"] - thr), 6)
        if out["method"] == "fallback_rule":
            out["uncertain"] = 1
        else:
            out["uncertain"] = int(margin is not None and out["margin"] <= float(margin) + 1e-12)
        pl = getattr(reg, "platt", {}).get(crit)
        out["p_cal"] = (round(float(sigmoid(pl[0] * out["score"] + pl[1])), 6)
                        if (pl is not None and out["method"] != "fallback_rule") else None)
        return out

    def fallback_rule(self, crit: str, feats: Dict[str, Any]) -> float:
        rule = self.cfg["fallback_rules"].get(crit)
        if not rule:
            return 0.0
        x = feats.get(rule["feature"])
        if x is None or not np.isfinite(x):
            return 0.0  # признак не измерен -> консервативно "нет нарушения"
        return sigmoid(float(rule.get("direction", 1)) * (float(x) - float(rule["center"])) / float(rule["scale"]))

    def any_violation_prob(self, region: str, crit_results: Dict[str, Dict[str, Any]],
                           feats: Dict[str, Any], embs: Optional[Dict[str, np.ndarray]]) -> float:
        reg = self.registry
        probs: List[float] = []
        emb = _emb_for(reg.any_emb.get(region), embs)
        # 1) отдельная модель "есть нарушение", если сохранена
        mb_g, mb_e = reg.any_geom.get(region), reg.any_emb.get(region)
        if mb_g is not None:
            try:
                cols = mb_g.meta.get("feature_cols") or sorted({c for cr in self.cfg["criteria_by_region"][region]
                                                                 for c in self.cfg["geometry_cols"][cr]})
                vec = np.asarray([float(feats.get(c) if feats.get(c) is not None and np.isfinite(feats.get(c))
                                        else reg.geometry_medians.get(c, 0.0)) for c in cols])
                probs.append(clip01(mb_g.predict(vec)))
            except Exception as e:  # noqa: BLE001
                LOG.warning("any-geom model failed: %s", e)
        if mb_e is not None and emb is not None:
            try:
                probs.append(clip01(mb_e.predict(emb)))
            except Exception as e:  # noqa: BLE001
                LOG.warning("any-emb model failed: %s", e)
        # 2) агрегация по критериям
        scores = [r["score"] for r in crit_results.values() if r.get("score") is not None]
        if self.cfg["stacking"].get("any_violation_aggregation", "max") == "noisy_or":
            crit_agg = clip01(1.0 - float(np.prod([1.0 - s for s in scores]))) if scores else None
        else:
            crit_agg = clip01(max(scores)) if scores else None
        any_model = float(np.mean(probs)) if probs else None
        if any_model is None and crit_agg is None:
            return float(self.cfg["output"]["fallback_quality_prob"])
        if any_model is None:
            return crit_agg
        if crit_agg is None:
            return any_model
        # 3) смесь двух оценок. OOF ROC-AUC по models/metrics_oof_full.json (by_region_binary.*.roc_auc_components,
        #    файл от 23.09.2026): any-модель 0.766/0.702 (spine/hip), смесь 0.5/0.5 0.823/0.759,
        #    после согласования с классом (consistent_quality_prob) 0.813/0.760; src/eval_oof_metrics.py.
        w = float(self.cfg["stacking"].get("any_blend_weight_model", 0.5))
        return clip01(w * any_model + (1.0 - w) * crit_agg)

    @staticmethod
    def consistent_quality_prob(prob: float, quality_class: int) -> float:
        """Согласование quality_prob с quality_class: класс определяется флагами
        критериев (порог подобран по OOF на каждый критерий), а вероятность —
        смесью моделей. Чтобы строка была непротиворечивой (class=1 <=> prob>=0.5)
        и ROC-AUC учитывал решение по критериям, вероятность монотонно сжимается
        в [0.5, 1] при нарушении и в [0, 0.5) при норме (порядок внутри класса
        сохраняется). OOF ROC-AUC смеси до/после согласования (models/metrics_oof_full.json,
        23.09.2026): spine 0.823 -> 0.813, hip 0.759 -> 0.760 — согласование нужно для
        непротиворечивости строки, а не для роста AUC."""
        p = clip01(prob)
        return 0.5 + 0.5 * p if quality_class else min(0.5 * p, 0.499999)

    # ---- обработка одного файла (никогда не бросает исключение наружу)
    def process_file(self, path: Path, root: Optional[Path] = None) -> Tuple[Dict[str, Any], Dict[str, Any]]:
        t0 = time.perf_counter()
        cfg_out = self.cfg["output"]
        path = Path(path)
        debug: Dict[str, Any] = {"file": str(path)}
        try:
            info = read_and_validate(path, self.cfg)
            # Область применения (src/region_support.py): чужой аппарат, неподдерживаемая область или
            # проекция, чужая модальность, геометрия кадра вне диапазона -> строка Failure с причиной
            # в debug CSV (region_supported = 0, region_support_reason). До классификации: вердикт
            # вне области применения не выдаётся. На 499 файлах заказчика не срабатывает.
            # Явное противоречие в тегах -> Failure и в пакетном пути, и в API. Геометрия кадра вне
            # наблюдаемого диапазона -> в пакетном пути строка обрабатывается как обычно (кадр того же
            # аппарата с другим размером не теряем на закрытом тесте), а отказ фиксируется в debug:
            # API и кабинет по нему выдают «вне области применения».
            reg_ok, reg_reason = region_support_tags_of(info)
            if reg_ok:
                geo_ok, geo_reason = region_support_geometry_of(info)
                if not geo_ok:
                    reg_ok, reg_reason = geo_ok, geo_reason
                    debug["region_support_scope"] = "geometry"
            else:
                debug["region_support_scope"] = "tags"
            debug["region_supported"] = int(reg_ok)
            debug["region_support_reason"] = reg_reason
            if not reg_ok and debug.get("region_support_scope") == "tags":
                raise UnsupportedInputError(reg_reason)
            region, region_src = classify_region(info, path, self.cfg)
            std_cols = set(int(c) for c in self.cfg["regions"].get("standard_cols", [300, 280, 248]))
            if (region_src.startswith("dims") and info.cols not in std_cols
                    and self.region_fallback is not None):
                # нестандартная ширина -> правило по ширине ненадёжно, решаем по содержимому
                emb0 = self.embedder.extract(info.img_u8)
                if emb0 is not None:
                    is_spine = int(self.region_fallback["model"].predict(emb0[None, :])[0]) == 1
                    if is_spine:
                        region, region_src = "spine", "content_fallback"
                    else:
                        region, side_src = detect_hip_side_region(info.img_u8)
                        region_src = f"content_fallback+{side_src}"
                    LOG.info("%s: non-standard width %d px -> region by content: %s", path.name, info.cols, region)
            debug["region_src"] = region_src
            debug.update(self._origin_hashes(path, info.ds))
            debug["sop_class_uid"] = _tag(info.ds, "SOPClassUID")
            debug["series_uid"] = _tag(info.ds, "SeriesInstanceUID")
            if self.sr_study and info.study_uid not in self._study_headers:
                try:
                    from dicom_sr import study_header_from_ds
                    self._study_headers[info.study_uid] = study_header_from_ds(info.ds)
                except Exception:  # noqa: BLE001
                    self._study_headers[info.study_uid] = {}
            # К11: признаки контура A для каждого варианта предобработки, который нужен критериям региона.
            # baseline считается всегда: на нём обучены any-модели, от него идут оверлеи, карточка и SR.
            crit_preproc = {c: preproc_variants(self.cfg, c) for c in self.cfg["criteria_by_region"][region]}
            geom_variants = {"baseline"} | {p["geom"] for p in crit_preproc.values()}
            feats_by_variant = {v: extract_geometry(info, region, self.cfg, variant=v) for v in sorted(geom_variants)}
            feats = feats_by_variant["baseline"]

            need_emb = any((region, c) in self.registry.emb for c in self.cfg["criteria_by_region"][region]) \
                or region in self.registry.any_emb
            mirror = False
            if need_emb and region == "right_hip":
                # унифицированная модель бедра может требовать зеркалирования правого
                mb = next((self.registry.emb.get((region, c)) for c in self.cfg["criteria_by_region"][region]
                           if self.registry.emb.get((region, c)) is not None), None)
                mirror = bool(mb.meta.get("mirror_right", False)) if mb is not None else False
            # К11: эмбеддинги по (вариант кадра, источник). any-модель всегда на baseline.
            embs_by_variant: Dict[str, Dict[str, np.ndarray]] = {}
            if need_emb:
                need: Dict[str, set] = {}
                for c in self.cfg["criteria_by_region"][region]:
                    mb = self.registry.emb.get((region, c))
                    if mb is None:
                        continue
                    src = str(mb.meta.get("emb_source", EmbeddingExtractor.DEFAULT_SOURCE))
                    need.setdefault(crit_preproc[c]["emb"], set()).add(src)
                mb_any = self.registry.any_emb.get(region)
                if mb_any is not None:
                    need.setdefault("baseline", set()).add(
                        str(mb_any.meta.get("emb_source", EmbeddingExtractor.DEFAULT_SOURCE)))
                for v, sources in need.items():
                    frame = info.img_u8 if (v == "baseline" or info.img_canonical is None) else info.img_canonical
                    embs_by_variant[v] = self.embedder.extract_many(frame, sources, mirror=mirror)
            # H2 (2.4.0): признаки контура A из эмбеддинга (sp_pos: synth_pos_logit) — до скоринга критериев
            self._add_embedding_features(region, info, crit_preproc, feats_by_variant, embs_by_variant, mirror)
            embs = embs_by_variant.get("baseline", {})
            emb = embs.get(EmbeddingExtractor.DEFAULT_SOURCE) if embs else None
            if self.extras:  # [EXTRAS] входы для extras (незеркалированный imagenet-эмбеддинг, как в data/embeddings.npy)
                try:
                    import extras as _extras
                    emb_x = emb if (emb is not None and not mirror) else self.embedder.extract(info.img_u8)
                    self._extras_input = {"img_u8": info.img_u8, "emb": emb_x,
                                          "tags": _extras.tags_from_dataset(info.ds),
                                          "pixel_hash": _extras.pixel_hash(info.ds)}
                except Exception as ex:  # noqa: BLE001  — extras не должны ломать основной путь
                    LOG.warning("extras input failed for %s: %s", path, ex)
                    self._extras_input = None

            crit_results = {c: self.score_criterion(
                region, c,
                feats_by_variant.get(crit_preproc[c]["geom"], feats),
                embs_by_variant.get(crit_preproc[c]["emb"], embs))
                for c in self.cfg["criteria_by_region"][region]}
            violations: List[str] = []
            for c, r in crit_results.items():
                if r["flag"]:
                    name = self.cfg["violations"][c]
                    if name not in violations:
                        violations.append(name)
            quality_class = 1 if violations else 0
            quality_prob = self.any_violation_prob(region, crit_results, feats, embs)
            debug["quality_prob_raw"] = round(float(quality_prob), 6)
            if self.cfg["stacking"].get("consistent_quality_prob", True):
                quality_prob = self.consistent_quality_prob(quality_prob, quality_class)
            # инвариант: quality_class == 1  <=>  quality_prob >= 0.5; класс и список нарушений согласованы всегда.

            row = {
                "path_to_study": self._path_str(path, root),
                "study_uid": info.study_uid,
                "image_uid": info.image_uid,
                "anatomical_region": official_region_name(region, self.cfg),
                "quality_class": int(quality_class),
                "violation_type": cfg_out["violation_separator"].join(violations),
                "quality_prob": round(float(quality_prob), 6),
                "processing_status": cfg_out["status_success"],
                "time_of_processing": 0.0,
            }
            debug.update({"internal_region": region, "region_source": region_src,
                          "rows": info.rows, "cols": info.cols, "warnings": "; ".join(info.warnings),
                          "embedding_used": bool(embs), "embedding_sources": "+".join(sorted(embs)) if embs else "",
                          **{f"feat_{k}": v for k, v in feats.items()
                                                                  if not isinstance(v, np.ndarray)}})
            # H2 (2.4.0): признак контура A из эмбеддинга лежит в варианте предобработки критерия (не baseline)
            for v_feats in feats_by_variant.values():
                if "synth_pos_logit" in v_feats:
                    debug["feat_synth_pos_logit"] = v_feats["synth_pos_logit"]
            for c, r in crit_results.items():
                for k in ("p_geom", "p_emb", "w_geom", "score", "threshold", "flag", "method", "margin", "uncertain", "p_cal"):
                    debug[f"{c}_{k}"] = r.get(k)
            # C1 (только debug/API): значения признаков, которые реально подаются в модель контура A
            # критерия (feature_cols модели, вариант предобработки критерия), и ранги контуров A/B.
            for c in crit_results:
                debug[f"{c}_model_features"] = _model_features_json(
                    self.registry.geom.get((region, c)), feats_by_variant.get(crit_preproc[c]["geom"], feats),
                    crit_preproc[c]["geom"])
                debug[f"{c}_rank_geom"] = crit_results[c].get("rank_geom")
                debug[f"{c}_rank_emb"] = crit_results[c].get("rank_emb")
            # К3: строка «не уверен», если не уверен хотя бы один критерий региона; risk_level — правило
            # calibration_utils.risk_level (высокий: class=1 и уверен; средний: не уверен; низкий: class=0 и уверен)
            needs_review = int(any(int(r.get("uncertain", 0)) for r in crit_results.values()))
            debug["needs_review"] = needs_review
            debug["risk_level"] = risk_level(quality_class, needs_review)
            debug["uncertain_criteria"] = ";".join(c for c, r in crit_results.items() if int(r.get("uncertain", 0)))
            if info.warnings:
                LOG.info("%s: %s", path.name, "; ".join(info.warnings))

            self._emit_bonus_outputs(path, info, region, feats, crit_results, quality_class,
                                      violations, quality_prob, debug)
        except Exception as e:  # noqa: BLE001  — ЖЕЛЕЗНОЕ правило: строка всегда есть
            err = f"{type(e).__name__}: {e}"
            LOG.error("FAILURE %s -> %s", path, err)
            LOG.debug(traceback.format_exc())
            region = guess_region_without_pixels(path, self.cfg)
            rel = self._path_str(path, root)
            study_uid, image_uid = self._uids_without_pixels(path, rel)
            row = {
                "path_to_study": rel,
                "study_uid": study_uid,
                "image_uid": image_uid,
                "anatomical_region": official_region_name(region, self.cfg),
                "quality_class": 0,
                "violation_type": "",
                "quality_prob": failure_quality_prob(cfg_out),
                "processing_status": cfg_out["status_failure"],
                "time_of_processing": 0.0,
            }
            debug.update({"internal_region": region, "region_source": "fallback", "error": err,
                          "needs_review": 1, "risk_level": risk_level(0, True), "uncertain_criteria": ""})
            debug.update(self._origin_hashes(path, None))
        row["time_of_processing"] = round(time.perf_counter() - t0, 4)
        debug["time_of_processing"] = row["time_of_processing"]
        return row, debug

    # ---- бонус-выходы (визуализация / DICOM SR / предложение коррекции области интереса) — опционально,
    # никогда не бросает исключение наружу (process_file уже внутри своего try, но эта
    # функция сама оборачивает каждый под-шаг отдельно, чтобы сбой одного бонуса не
    # тушил остальные и уж тем более не портил основную строку CSV).
    def _emit_bonus_outputs(self, path: Path, info, region: str, feats: Dict[str, Any],
                            crit_results: Dict[str, Dict[str, Any]], quality_class: int,
                            violations: List[str], quality_prob: float, debug: Dict[str, Any]) -> None:
        # sc_dir в списке (2.4.1): без него --sc-dir без других бонус-флагов не писал Secondary Capture
        if not (self.visualize_dir or self.sr_dir or self.roi_autocorrect_dir or self.seg_dir or self.sc_dir):
            return
        stem = self._bonus_stem(path)
        violation_type_str = self.cfg["output"]["violation_separator"].join(violations)

        if self.visualize_dir or self.sc_dir:
            try:
                import cv2 as _cv2
                from visualize_report import render_overlay, save_overlay_sc
                overlay = render_overlay(info.img_u8, region, feats, crit_results,
                                         quality_class, violation_type_str)
                if self.visualize_dir:
                    out_png = self.visualize_dir / f"{stem}_overlay.png"
                    _cv2.imwrite(str(out_png), overlay)
                    debug["bonus_overlay_png"] = str(out_png)
                if self.sc_dir:
                    # бонус ТЗ 2.6: та же картинка как DICOM Secondary Capture рядом с исходной серией
                    out_sc = self.sc_dir / f"{stem}_overlay.dcm"
                    save_overlay_sc(str(out_sc), overlay, info.ds,
                                    model_version=__version__, config_hash=config_hash(self.cfg))
                    debug["bonus_overlay_dcm"] = str(out_sc)
            except Exception as e:  # noqa: BLE001
                LOG.warning("visualize_report failed for %s: %s", path.name, e)
                debug["bonus_overlay_error"] = str(e)

        if self.sr_dir and (self.sr_per_image or not self.sr_study):
            try:
                from dicom_sr import save_sr
                out_sr = self.sr_dir / f"{stem}_sr.dcm"
                save_sr(str(out_sr), info.ds, region, quality_class, violation_type_str,
                        quality_prob, feats)
                debug["bonus_sr_dcm"] = str(out_sr)
            except Exception as e:  # noqa: BLE001
                LOG.warning("dicom_sr failed for %s: %s", path.name, e)
                debug["bonus_sr_error"] = str(e)

        if self.roi_autocorrect_dir and region in ("right_hip", "left_hip"):
            try:
                from auto_roi import suggest_hip_roi, draw_roi_correction
                suggestion = suggest_hip_roi(info.img_u8, feats)
                debug["bonus_roi_needs_correction"] = suggestion.get("needs_correction")
                debug["bonus_roi_reason"] = suggestion.get("reason")
                debug["bonus_roi_deficit_mm"] = suggestion.get("deficit_mm")
                debug.update(roi_suggestion_debug_fields(suggestion, info))  # идея 23: рамка и шаг пикселя для API
                if suggestion.get("needs_correction"):
                    import cv2
                    out_png = self.roi_autocorrect_dir / f"{stem}_roi_correction.png"
                    overlay = draw_roi_correction(info.img_u8, suggestion)
                    cv2.imwrite(str(out_png), overlay)
                    debug["bonus_roi_png"] = str(out_png)
            except Exception as e:  # noqa: BLE001
                LOG.warning("auto_roi failed for %s: %s", path.name, e)
                debug["bonus_roi_error"] = str(e)

        if self.seg_dir:
            # бонус «сегментация»: маски структур (кость, посторонние предметы / поле сканирования, область
            # интереса) как DICOM SEG + PNG + JSON; только для Success-строк (сюда Failure не доходит)
            try:
                from segmentation_export import export_segmentation
                seg = export_segmentation(info.img_u8, region, info.ds, str(self.seg_dir), stem,
                                          image_uid=info.image_uid, study_uid=info.study_uid,
                                          model_version=__version__, config_hash=config_hash(self.cfg),
                                          source_path=str(path))
                debug["bonus_seg_dcm"] = seg["seg_dcm"]
                debug["bonus_seg_png"] = seg["png"]
                debug["bonus_seg_json"] = seg["json"]
                debug["bonus_seg_n_segments"] = seg["n_segments"]
                debug["bonus_seg_bone_area_px"] = seg["areas_px"].get("bone", 0)
            except Exception as e:  # noqa: BLE001
                LOG.warning("segmentation_export failed for %s: %s", path.name, e)
                debug["bonus_seg_error"] = str(e)

    def _bonus_stem(self, path: Path) -> str:
        """Имя бонус-файла, уникальное в пределах прогона.

        В DXA-датасетах имена файлов повторяются (`CR000001.dcm` есть почти в каждом
        исследовании), поэтому `path.stem` затирал бы оверлеи и SR разных исследований в
        одной папке. Берём имя файла и короткий хэш полного пути: имя остаётся читаемым,
        а совпадений нет. Соответствие «бонус-файл -> исходный снимок» есть в debug-CSV
        (колонки `bonus_overlay_png`, `bonus_sr_dcm`, `bonus_roi_png`).
        """
        h = hashlib.sha1(str(path.resolve()).encode("utf-8", errors="ignore")).hexdigest()[:10]
        safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", path.stem)[:60] or "image"
        return f"{safe}__{h}"

    def _path_str(self, path: Path, root: Optional[Path]) -> str:
        mode = self.cfg["output"].get("path_mode", "relative")
        try:
            if mode == "name":
                return path.name
            if mode == "absolute" or root is None:
                return str(path.resolve())
            shown = getattr(self, "_display_paths", {}).get(path.resolve())
            if shown:
                return shown
            return str(path.resolve().relative_to(Path(root).resolve()))
        except Exception:  # noqa: BLE001
            return str(path)

    @staticmethod
    def _uids_without_pixels(path: Path, rel: Optional[str] = None) -> Tuple[str, str]:
        """UID для строки Failure: из тегов, если читаются; иначе детерминированный хэш
        СОДЕРЖИМОГО файла и папки — одинаковый на любой машине, в контейнере и при переименовании
        (закрытый набор приходит с другими именами файлов)."""
        try:
            ds = pydicom.dcmread(str(path), force=True, stop_before_pixels=True)
            s, i = _tag(ds, "StudyInstanceUID"), _tag(ds, "SOPInstanceUID")
        except Exception:  # noqa: BLE001
            s, i = "", ""
        try:
            fallback_study = dir_content_hash(path.parent)
            fallback_image = content_hash(path)
        except Exception:  # noqa: BLE001
            rel_p = Path(rel) if rel else path
            fallback_study, fallback_image = path_hash(rel_p.parent), path_hash(rel_p)
        return (s or "hash-" + fallback_study), (i or "hash-" + fallback_image)

    @staticmethod
    def _origin_hashes(path: Path, ds) -> Dict[str, str]:
        """SHA-256 исходного файла (побайтово) и массива пикселей — хэш неизменности оригинала.
        Никогда не бросает исключение."""
        out = {"sha256_file": "", "sha256_pixels": ""}
        try:
            h = hashlib.sha256()
            with open(path, "rb") as fh:
                for chunk in iter(lambda: fh.read(1 << 20), b""):
                    h.update(chunk)
            out["sha256_file"] = h.hexdigest()
        except Exception:  # noqa: BLE001
            pass
        try:
            if ds is not None and "PixelData" in ds:
                out["sha256_pixels"] = hashlib.sha256(bytes(ds.PixelData)).hexdigest()
        except Exception:  # noqa: BLE001
            pass
        return out

    def write_study_sr(self, rows: List[Dict[str, Any]], debug_rows: List[Dict[str, Any]],
                       out_dir: Path) -> Dict[str, str]:
        """Один DICOM SR на каждое исследование партии (включая норму и Failure).
        Возвращает {study_uid: путь к SR}. Ошибка одного исследования не мешает остальным."""
        from dicom_sr import save_study_sr, study_sr_filename
        out_dir = Path(out_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        sep = self.cfg["output"]["violation_separator"]
        model_version = str(self.cfg.get("version", __version__))
        cfg_hash = config_hash(self.cfg)
        groups: Dict[str, List[Dict[str, Any]]] = {}
        for i, r in enumerate(rows):
            dbg = debug_rows[i] if i < len(debug_rows) else {}
            groups.setdefault(str(r.get("study_uid") or ""), []).append({
                "image_uid": str(r.get("image_uid") or ""),
                "sop_class_uid": dbg.get("sop_class_uid") or "",
                "series_uid": dbg.get("series_uid") or "",
                "anatomical_region": str(r.get("anatomical_region") or ""),
                "quality_class": int(r.get("quality_class") or 0),
                "violations": [v.strip() for v in str(r.get("violation_type") or "").split(sep) if v.strip()],
                "quality_prob": r.get("quality_prob"),
                "processing_status": str(r.get("processing_status") or ""),
                "sha256_file": dbg.get("sha256_file") or "",
                "sha256_pixels": dbg.get("sha256_pixels") or "",
                "path_to_study": str(r.get("path_to_study") or ""),
                # C1: для «Приоритетного действия» (на 9 колонок и дайджест SOP UID не влияет)
                "internal_region": dbg.get("internal_region") or "",
                "criteria": _criteria_for_priority(self.cfg, dbg),
                "roi_route": _roi_route_for_priority(self.cfg, dbg),
                # 2.5: основание команды и второе мнение по оси (текст в SR; класс и флаги не меняются)
                **_evidence_texts_for_sr(self.cfg, dbg, str(r.get("processing_status") or "")),
            })
        written: Dict[str, str] = {}
        for study_uid, items in groups.items():
            if not study_uid:
                continue
            try:
                out = out_dir / study_sr_filename(study_uid)
                save_study_sr(str(out), study_uid, items, model_version, cfg_hash,
                              self._study_headers.get(study_uid, {}))
                written[study_uid] = str(out)
            except Exception as e:  # noqa: BLE001
                LOG.warning("study SR failed for %s: %s", study_uid, e)
        LOG.info("Study SR written: %d of %d studies -> %s", len(written), len([g for g in groups if g]), out_dir)
        return written

    # ---- пакет
    def run(self, input_path: Path, output_csv: Path, debug_csv: Optional[Path] = None,
            xlsx: bool = False, limit: Optional[int] = None,
            keep_temp: Optional[List[Path]] = None) -> List[Dict[str, Any]]:
        """Пакетная обработка. После вызова `self.last_row_files[i]` — фактический файл на диске, по
        которому получена строка i (для файлов из zip — путь во временном каталоге распаковки).

        keep_temp (2.4.1, Б6): если передан список, временные каталоги распаковки zip (`densito_in_*`)
        не удаляются, а добавляются в этот список — вызывающий (API) читает по `last_row_files` теги,
        кадры и заголовки и удаляет каталоги сам в finally. По умолчанию (None, пакетный режим CLI)
        каталоги удаляются здесь же, как раньше."""
        tmp_dirs: List[Path] = []
        t_start = time.perf_counter()
        rows: List[Dict[str, Any]] = []
        debug_rows: List[Dict[str, Any]] = []
        row_files: List[str] = []
        self.last_row_files = row_files
        try:
            self._display_paths: Dict[Path, str] = {}
            root, files = discover_files(Path(input_path), tmp_dirs, self._display_paths)
            if limit:
                files = files[:limit]
            LOG.info("Found %d DICOM candidate files under %s", len(files), input_path)
            if not files:
                LOG.warning("No DICOM files found -> writing empty CSV with header")
            if files and (self.registry.emb or self.registry.any_emb):
                # прогрев всех нужных backbone вне замера time_of_processing
                for mb in list(self.registry.emb.values()) + list(self.registry.any_emb.values()):
                    self.embedder._init(str(mb.meta.get("emb_source", EmbeddingExtractor.DEFAULT_SOURCE)))
            extras_inputs: List[Optional[Dict[str, Any]]] = []  # [EXTRAS]
            for i, f in enumerate(files, 1):
                self._extras_input = None
                row, dbg = self.process_file(f, root)
                rows.append(row)
                debug_rows.append(dbg)
                row_files.append(str(f))
                extras_inputs.append(self._extras_input)
                if i % 25 == 0 or i == len(files):
                    LOG.info("  %d/%d processed (%.1fs)", i, len(files), time.perf_counter() - t_start)
        finally:
            if keep_temp is not None:
                keep_temp.extend(tmp_dirs)      # удалит вызывающий (API) после чтения файлов строк
            else:
                for d in tmp_dirs:
                    shutil.rmtree(d, ignore_errors=True)

        write_results(rows, Path(output_csv), self.cfg, xlsx=xlsx)
        self.last_debug_rows = debug_rows  # для API: детали по критериям без повторного чтения CSV
        try:
            self.last_study_completeness = study_completeness_by_study(rows)
        except Exception as e:  # noqa: BLE001 — примечание не должно ломать основной выход
            LOG.warning("study completeness not computed: %s", e)
            self.last_study_completeness = {}
        self.last_study_sr = {}
        if self.sr_study:
            try:
                sr_dir = self.sr_study_dir or (Path(output_csv).with_suffix(".csv").parent / "sr")
                self.last_study_sr = self.write_study_sr(rows, debug_rows, sr_dir)
            except Exception as e:  # noqa: BLE001  — бонус не должен ломать основной выход
                LOG.warning("study SR not written: %s", e)
        if debug_csv:
            write_debug(debug_rows, Path(debug_csv))
        self.last_extras_rows = []
        if self.extras:  # [EXTRAS] отдельный файл <output>_extras.csv; основной CSV уже записан
            try:
                import extras as _extras
                self.last_extras_rows = _extras.compute_extras_for_rows(rows, debug_rows, extras_inputs,
                                                                        models_dir=self.registry.models_dir)
                extras_csv = Path(output_csv).with_suffix(".csv")
                extras_csv = extras_csv.with_name(extras_csv.stem + "_extras.csv")
                _extras.write_extras_csv(self.last_extras_rows, extras_csv)
                LOG.info("extras written -> %s", extras_csv)
            except Exception as e:  # noqa: BLE001  — extras не должны ломать основной выход
                LOG.warning("extras not written: %s", e)
        n_fail = sum(1 for r in rows if r["processing_status"] == self.cfg["output"]["status_failure"])
        LOG.info("DONE: %d rows, %d failures, %.1fs total -> %s", len(rows), n_fail,
                 time.perf_counter() - t_start, output_csv)
        return rows


# --------------------------------------------------------------------------- #
# Запись результатов
# --------------------------------------------------------------------------- #
def study_completeness_by_study(rows: List[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    """{study_uid: {"spine": bool, "hip": bool, "note": str|None}} по строкам официального формата.

    Та же функция dicom_sr.study_completeness, что формирует примечание в SR исследования, поэтому
    JSON API и SR согласованы по построению. Официальные колонки не читаются на запись и не меняются."""
    from dicom_sr import study_completeness
    groups: Dict[str, List[Dict[str, Any]]] = {}
    for r in rows:
        uid = str(r.get("study_uid") or "")
        if uid:
            groups.setdefault(uid, []).append({"anatomical_region": str(r.get("anatomical_region") or ""),
                                               "processing_status": str(r.get("processing_status") or "")})
    return {uid: study_completeness(items) for uid, items in groups.items()}


def write_results(rows: List[Dict[str, Any]], output_csv: Path, cfg: Dict[str, Any], xlsx: bool = False):
    cols = cfg["output"]["columns"]
    output_csv = Path(output_csv)
    if output_csv.suffix.lower() == ".xlsx":       # запросили xlsx -> CSV пишем рядом, xlsx по указанному пути
        xlsx = True
        output_csv = output_csv.with_suffix(".csv")
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    tmp = output_csv.with_suffix(output_csv.suffix + ".tmp")
    with open(tmp, "w", newline="", encoding=cfg["output"].get("csv_encoding", "utf-8")) as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore", quoting=csv.QUOTE_MINIMAL,
                           lineterminator="\n")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in cols})
    os.replace(tmp, output_csv)  # атомарная запись
    if xlsx:
        try:
            xlsx_path = output_csv.with_suffix(".xlsx")
            write_xlsx(rows, xlsx_path, cfg)
            LOG.info("XLSX written: %s", xlsx_path)
        except Exception as e:  # noqa: BLE001
            LOG.warning("XLSX not written (%s); CSV is the primary artifact", e)


def write_xlsx(rows: List[Dict[str, Any]], xlsx_path: Path, cfg: Dict[str, Any]) -> None:
    """XLSX для человека: лист «Результаты» с теми же колонками, что и CSV (закреплённая
    шапка, автофильтр, подсветка нарушений и сбоев), лист «Сводка» (счётчики по областям и
    типам нарушений) и лист «Исследования» (агрегат по study_uid). Официальный артефакт —
    CSV; XLSX — удобное представление тех же данных."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols = cfg["output"]["columns"]
    fail = cfg["output"]["status_failure"]
    sep = cfg["output"]["violation_separator"]
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="01696F")
    fill_viol = PatternFill("solid", fgColor="FBE9F2")
    fill_fail = PatternFill("solid", fgColor="F3E4D5")
    fill_ok = PatternFill("solid", fgColor="E9F1E3")

    def _header(ws, names):
        ws.append(list(names))
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(names))}{max(2, ws.max_row)}"

    def _autowidth(ws, minimum=8, maximum=60):
        for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 500)), 1):
            w = max((len(str(c.value)) for c in col if c.value is not None), default=minimum)
            ws.column_dimensions[get_column_letter(i)].width = max(minimum, min(maximum, w + 2))

    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"
    extra = ["Действие"]
    _header(ws, cols + extra)
    for r in rows:
        is_fail = r.get("processing_status") == fail
        is_viol = str(r.get("quality_class")) == "1"
        action = "Проверить вручную" if is_fail else ("Проверить / переснять" if is_viol else "Принять")
        ws.append([r.get(c, "") for c in cols] + [action])
        f = fill_fail if is_fail else (fill_viol if is_viol else fill_ok)
        for c in ws[ws.max_row]:
            c.fill = f
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols) + len(extra))}{max(2, ws.max_row)}"
    _autowidth(ws)

    # --- Сводка
    ws2 = wb.create_sheet("Сводка")
    n = len(rows)
    n_fail = sum(1 for r in rows if r.get("processing_status") == fail)
    n_viol = sum(1 for r in rows if str(r.get("quality_class")) == "1")
    n_ok = n - n_fail - n_viol
    t_sum = sum(float(r.get("time_of_processing") or 0) for r in rows)
    studies = {r.get("study_uid") for r in rows if r.get("study_uid")}
    _header(ws2, ["Показатель", "Значение"])
    for k, v in [
        ("Файлов обработано", n), ("Исследований", len(studies)),
        ("Без нарушений", n_ok), ("С нарушениями", n_viol),
        ("Доля с нарушениями", f"{(n_viol / n * 100):.1f} %" if n else "—"),
        ("Ошибок обработки (Failure)", n_fail),
        ("Суммарное время обработки, с", round(t_sum, 2)),
        ("Среднее время на файл, с", round(t_sum / n, 3) if n else "—"),
        ("Версия модели", cfg.get("version", __version__)),
    ]:
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["Область", "Файлов", "С нарушениями", "Ошибок"])
    for c in ws2[ws2.max_row]:
        c.font = Font(bold=True)
    for reg in sorted({r.get("anatomical_region", "") for r in rows}):
        sub = [r for r in rows if r.get("anatomical_region", "") == reg]
        ws2.append([reg or "—", len(sub), sum(1 for r in sub if str(r.get("quality_class")) == "1"),
                    sum(1 for r in sub if r.get("processing_status") == fail)])
    ws2.append([])
    ws2.append(["Тип нарушения", "Файлов"])
    for c in ws2[ws2.max_row]:
        c.font = Font(bold=True)
    counts: Dict[str, int] = {}
    for r in rows:
        for v in str(r.get("violation_type") or "").split(sep):
            if v.strip():
                counts[v.strip()] = counts.get(v.strip(), 0) + 1
    for k, v in sorted(counts.items(), key=lambda kv: -kv[1]):
        ws2.append([k, v])
    ws2.auto_filter.ref = None
    _autowidth(ws2, minimum=14)

    # --- Исследования
    ws3 = wb.create_sheet("Исследования")
    _header(ws3, ["study_uid", "Файлов", "С нарушениями", "Ошибок", "Макс. quality_prob", "Нарушения"])
    by_study: Dict[str, List[Dict[str, Any]]] = {}
    for r in rows:
        by_study.setdefault(str(r.get("study_uid") or ""), []).append(r)
    for uid, sub in sorted(by_study.items(), key=lambda kv: -max(float(x.get("quality_prob") or 0) for x in kv[1])):
        viols: List[str] = []
        for r in sub:
            for v in str(r.get("violation_type") or "").split(sep):
                if v.strip() and v.strip() not in viols:
                    viols.append(v.strip())
        nv = sum(1 for r in sub if str(r.get("quality_class")) == "1")
        ws3.append([uid, len(sub), nv, sum(1 for r in sub if r.get("processing_status") == fail),
                    # 6 знаков, как в CSV: при 3 знаках 0.499999 (Failure/норма) показывалось бы как 0.5
                    round(max(float(x.get("quality_prob") or 0) for x in sub), 6), "; ".join(viols)])
        if nv:
            for c in ws3[ws3.max_row]:
                c.fill = fill_viol
    ws3.auto_filter.ref = f"A1:F{max(2, ws3.max_row)}"
    _autowidth(ws3, minimum=10)
    wb.save(xlsx_path)


# --------------------------------------------------------------------------- #
# Идея 23: предложение коррекции области интереса бедра для подтверждения специалистом.
# Только дополнительные колонки debug-CSV (bonus_roi_*); официальные 9 колонок не затрагиваются.
# --------------------------------------------------------------------------- #
def _box_to_str(box) -> str:
    """(x0, y0, x1, y1) -> "x0,y0,x1,y1" (пустая строка, если рамки нет)."""
    if not box:
        return ""
    try:
        return ",".join(str(int(round(float(v)))) for v in box)
    except (TypeError, ValueError):
        return ""


def _criteria_for_priority(cfg: Dict[str, Any], dbg: Dict[str, Any]) -> List[Dict[str, Any]]:
    """C1: критерии снимка из debug-словаря для dicom_sr.study_priority_action."""
    reg = str((dbg or {}).get("internal_region") or "")
    out = []
    for c in cfg.get("criteria_by_region", {}).get(reg, []):
        if dbg.get(f"{c}_flag") is None:
            continue
        out.append({"code": c, "score": dbg.get(f"{c}_score"), "threshold": dbg.get(f"{c}_threshold"),
                    "flag": dbg.get(f"{c}_flag"), "uncertain": dbg.get(f"{c}_uncertain")})
    # 2.5: основание команды по флагу (src/action_evidence.py) — только для действия на визит, класс не меняется
    try:
        import action_evidence as _ae
        ev = _ae.criterion_evidence_map(_ae.evidence_from_debug(dbg, cfg))
        for it in out:
            if it["code"] in ev:
                it["evidence_command"] = ev[it["code"]]["command"]
    except Exception:  # noqa: BLE001
        pass
    return out


def _evidence_texts_for_sr(cfg: Dict[str, Any], dbg: Dict[str, Any], status: str) -> Dict[str, str]:
    """2.5: {action_evidence_text, axis_second_opinion_text} для SR по debug-словарю снимка; пусто при Failure."""
    out = {"action_evidence_text": "", "axis_second_opinion_text": ""}
    if status.lower() == "failure":
        return out
    try:
        import action_evidence as _ae
        out["action_evidence_text"] = _ae.sr_text(_ae.evidence_from_debug(dbg, cfg))
        ax = _ae.axis_from_debug(dbg)
        out["axis_second_opinion_text"] = (ax or {}).get("text", "")
    except Exception:  # noqa: BLE001
        pass
    return out


def _roi_route_for_priority(cfg: Dict[str, Any], dbg: Dict[str, Any]) -> str:
    """C1: код развилки hip_roi (extras.hip_roi_reason) для снимка бедра с флагом *_roi, иначе ''."""
    reg = str((dbg or {}).get("internal_region") or "")
    if reg not in ("right_hip", "left_hip"):
        return ""
    crit = "rh_roi" if reg == "right_hip" else "lh_roi"
    try:
        if int(dbg.get(f"{crit}_flag") or 0) != 1:
            return ""
        import extras as _ex
        rr = _ex.hip_roi_reason(dbg, True)
        return rr["code"] if rr else ""
    except Exception:  # noqa: BLE001
        return ""


def _model_features_json(mb, feats: Dict[str, Any], variant: str) -> str:
    """C1: {"variant": ..., "values": {признак: значение}} по feature_cols модели контура A (JSON-строка).
    Пустая строка, если модели нет. Значение None — признак не посчитан (модель подставила медиану)."""
    if mb is None:
        return ""
    try:
        cols = list(mb.meta.get("feature_cols") or [])
        vals = {}
        for c in cols:
            v = feats.get(c)
            try:
                v = float(v)
                v = round(v, 4) if np.isfinite(v) else None
            except (TypeError, ValueError):
                v = None
            vals[c] = v
        return json.dumps({"variant": str(variant), "values": vals}, ensure_ascii=False)
    except Exception:  # noqa: BLE001
        return ""


def roi_suggestion_debug_fields(suggestion: Dict[str, Any], info) -> Dict[str, Any]:
    """Дополнительные debug-поля по предложению области интереса (auto_roi.suggest_hip_roi):
    рамка в пикселях исходного кадра, расширенная рамка (если скан короткий — нижняя граница
    может выходить за кадр), сторона и шаг пикселя (мм), чтобы API мог показать рамку в мм.
    Ничего не бросает: при любой ошибке возвращает пустые строки."""
    out: Dict[str, Any] = {}
    try:
        out["bonus_roi_box_px"] = _box_to_str(suggestion.get("suggested_box_px"))
        out["bonus_roi_ext_box_px"] = _box_to_str(suggestion.get("extended_box_px"))
        out["bonus_roi_bone_box_px"] = _box_to_str(suggestion.get("bone_box_px"))
        out["bonus_roi_side"] = str(suggestion.get("side_detected") or "")
        ps = getattr(info, "pixel_spacing", None)
        out["bonus_roi_pixel_spacing_mm"] = f"{float(ps[0]):.4f},{float(ps[1]):.4f}" if ps else ""
    except Exception:  # noqa: BLE001
        for k in ("bonus_roi_box_px", "bonus_roi_ext_box_px", "bonus_roi_bone_box_px",
                  "bonus_roi_side", "bonus_roi_pixel_spacing_mm"):
            out.setdefault(k, "")
    return out


def write_debug(debug_rows: List[Dict[str, Any]], path: Path):
    try:
        import pandas as pd
        path.parent.mkdir(parents=True, exist_ok=True)
        pd.DataFrame(debug_rows).to_csv(path, index=False)
        LOG.info("Debug CSV written: %s", path)
    except Exception as e:  # noqa: BLE001
        LOG.warning("Debug CSV not written: %s", e)


def allowed_violations_by_region(cfg: Dict[str, Any]) -> Dict[str, set]:
    """{официальная строка области: допустимые нарушения}. Источник — словарь организаторов
    schema/official_dictionary.json (violation_type по области); если файла нет — то же из
    config.yaml (criteria_by_region + violations). «Некорректная укладка» законна для обеих областей."""
    p = PROJECT_ROOT / "schema" / "official_dictionary.json"
    try:
        d = json.loads(p.read_text(encoding="utf-8"))
        return {str(k): set(v) for k, v in (d.get("violation_type") or {}).items()}
    except Exception as e:  # noqa: BLE001
        LOG.warning("official dictionary %s not read (%s): violations by region from config", p, e)
    out: Dict[str, set] = {}
    viol = cfg["violations"]
    for internal, crits in cfg["criteria_by_region"].items():
        name = cfg["regions"]["spine"] if internal == "spine" else cfg["regions"]["hip"]
        out.setdefault(name, set()).update(viol[c] for c in crits if c in viol)
    return out


def validate_output_csv(path: Path, cfg: Optional[Dict[str, Any]] = None,
                        schema_check: bool = True) -> List[str]:
    """Проверка выходного файла на соответствие официальному формату (правила из config.yaml
    + JSON Schema schema/results_row.schema.json). Возвращает список проблем."""
    cfg = cfg or load_config()
    problems: List[str] = []
    rows_for_schema: List[Dict[str, str]] = []
    allowed_regions = {cfg["regions"]["spine"], cfg["regions"]["hip"]}
    allowed_viol = set(cfg["violations"].values())
    viol_by_region = allowed_violations_by_region(cfg)
    sep = cfg["output"]["violation_separator"]
    with open(path, "r", encoding=cfg["output"].get("csv_encoding", "utf-8"), newline="") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != cfg["output"]["columns"]:
            problems.append(f"columns mismatch: {reader.fieldnames}")
        for i, r in enumerate(reader, 2):
            if r["anatomical_region"] not in allowed_regions:
                problems.append(f"line {i}: bad anatomical_region '{r['anatomical_region']}'")
            if r["quality_class"] not in ("0", "1"):
                problems.append(f"line {i}: quality_class must be 0/1, got '{r['quality_class']}'")
            try:
                p = float(r["quality_prob"])
                if not 0.0 <= p <= 1.0:
                    problems.append(f"line {i}: quality_prob out of [0,1]: {p}")
                # инвариант строки (в т. ч. Failure): quality_class 1 <=> quality_prob >= 0.5
                elif r["quality_class"] == "1" and p < 0.5:
                    problems.append(f"line {i}: quality_class=1 but quality_prob {p} < 0.5")
                elif r["quality_class"] == "0" and p >= 0.5:
                    problems.append(f"line {i}: quality_class=0 but quality_prob {p} >= 0.5")
            except ValueError:
                problems.append(f"line {i}: quality_prob not float: '{r['quality_prob']}'")
            if r["processing_status"] not in (cfg["output"]["status_success"], cfg["output"]["status_failure"]):
                problems.append(f"line {i}: bad processing_status '{r['processing_status']}'")
            if r["violation_type"]:
                region_viol = viol_by_region.get(r["anatomical_region"])
                for v in r["violation_type"].split(sep):
                    if v.strip() not in allowed_viol:
                        problems.append(f"line {i}: unknown violation '{v}'")
                    elif region_viol is not None and v.strip() not in region_viol:
                        problems.append(f"line {i}: violation '{v}' not allowed for region "
                                        f"'{r['anatomical_region']}'")
                if r["quality_class"] != "1":
                    problems.append(f"line {i}: violations listed but quality_class != 1")
            elif r["quality_class"] == "1":
                problems.append(f"line {i}: quality_class=1 but violation_type empty")
            try:
                float(r["time_of_processing"])
            except ValueError:
                problems.append(f"line {i}: time_of_processing not float")
            rows_for_schema.append(r)
    # JSON Schema (schema/results_row.schema.json) — второй, независимый источник истины формата.
    # Проверка через пакет jsonschema, если он установлен, иначе встроенный мини-валидатор
    # (src/schema_check.py). Отсутствие файла схемы — предупреждение, не ошибка формата.
    if schema_check:
        try:
            from schema_check import load_schema, validate_results_rows
            problems.extend(validate_results_rows(rows_for_schema, load_schema("results_row")))
        except FileNotFoundError as e:
            LOG.warning("JSON Schema not found, schema check skipped: %s", e)
        except Exception as e:  # noqa: BLE001
            LOG.warning("JSON Schema check failed to run: %s", e)
    return problems


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def setup_logging(log_file: Optional[Path], verbose: bool):
    handlers: List[logging.Handler] = [logging.StreamHandler(sys.stdout)]
    if log_file:
        log_file.parent.mkdir(parents=True, exist_ok=True)
        handlers.append(logging.FileHandler(log_file, encoding="utf-8"))
    logging.basicConfig(level=logging.DEBUG if verbose else logging.INFO,
                        format="%(asctime)s %(levelname)s %(message)s", handlers=handlers, force=True)


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="DensitoAI — пакетный инференс качества DXA (DICOM -> CSV)")
    ap.add_argument("--input", "-i", default=None, help="Папка с исследованиями, один DICOM или zip-архив (обязателен, кроме --validate-only)")
    ap.add_argument("--output", "-o", required=True, help="Путь к выходному CSV (или .xlsx)")
    ap.add_argument("--models-dir", default=None, help=f"Папка с моделями (по умолчанию {MODELS_DIR})")
    ap.add_argument("--config", default=None, help=f"config.yaml (по умолчанию {CONFIG_PATH})")
    ap.add_argument("--debug-csv", nargs="?", const="auto", default=None,
                    help="Расширенный CSV с признаками/скорами (по умолчанию <output>_debug.csv)")
    ap.add_argument("--xlsx", action="store_true", help="Дополнительно записать .xlsx рядом с CSV")
    ap.add_argument("--no-embeddings", action="store_true", help="Отключить контур B (только геометрия)")
    ap.add_argument("--path-mode", choices=["relative", "absolute", "name"], default=None)
    ap.add_argument("--limit", type=int, default=None, help="Обработать только первые N файлов (отладка)")
    ap.add_argument("--log-file", default=None, help="Файл лога (по умолчанию <output>.log)")
    ap.add_argument("--verbose", "-v", action="store_true")
    ap.add_argument("--validate-only", action="store_true",
                    help="Только проверить формат уже существующего CSV (--output) и выйти")
    ap.add_argument("--visualize-dir", default=None, help="[BONUS] Papka dlya PNG-overlayev s geometricheskimi priznakami kachestva")
    ap.add_argument("--sr-dir", default=None, help="[BONUS] Papka dlya DICOM Structured Report (.dcm) s rezultatami")
    ap.add_argument("--roi-autocorrect-dir", default=None, help="[BONUS] Papka dlya diagnosticheskih PNG s predlozheniem korrektsii ROI (bedro)")
    ap.add_argument("--sr-study", action="store_true",
                    help="[BONUS] Один DICOM SR на исследование (включая норму) в <каталог CSV>/sr/<study_uid>_SR.dcm")
    ap.add_argument("--sr-study-dir", default=None, help="[BONUS] Каталог для SR на исследование (включает --sr-study)")
    ap.add_argument("--sr-per-image", action="store_true",
                    help="[BONUS] Писать также SR на каждый снимок (по умолчанию только SR на исследование)")
    ap.add_argument("--sc-dir", default=None,
                    help="[BONUS] Каталог для DICOM Secondary Capture с визуализацией (серия с визуализацией, ТЗ 2.6)")
    ap.add_argument("--extras", action="store_true",
                    help="[EXTRAS] Дополнительно записать <output>_extras.csv (белые линии, OOD-gate, эндопротез, когерентность)")
    ap.add_argument("--seg-dir", default=None,
                    help="[BONUS] Каталог для экспорта сегментации структур: DICOM SEG (.dcm) + PNG-маска + JSON с контурами "
                         "на каждый Success-снимок (по умолчанию выключено, на results.csv не влияет)")
    ap.add_argument("--series-zip", default=None, metavar="PATH",
                    help="Архив дополнительных DICOM-серий (ТЗ п. 2.7): SR на исследование, наложение (SC), "
                         "сегментация (SEG). По умолчанию <каталог CSV>/additional_series.zip; отключается "
                         "переменной окружения DENSITO_SERIES_ZIP=0. На results.csv не влияет.")
    args = ap.parse_args(argv)

    output_csv = Path(args.output)
    if output_csv.suffix.lower() == ".xlsx":   # CSV — основной артефакт, xlsx пишется рядом
        args.xlsx = True
        output_csv = output_csv.with_suffix(".csv")
    setup_logging(Path(args.log_file) if args.log_file else output_csv.with_suffix(".log"), args.verbose)
    try:
        cfg = load_config(Path(args.config) if args.config else None)
    except ConfigError as e:
        LOG.critical("%s", e)
        print(f"Ошибка конфигурации: {e}", file=sys.stderr)
        return 2
    if args.path_mode:
        cfg["output"]["path_mode"] = args.path_mode

    if not args.validate_only and not args.input:
        ap.error("--input/-i обязателен (кроме режима --validate-only)")
    if args.validate_only:
        problems = validate_output_csv(output_csv, cfg)
        print("\n".join(problems) if problems else "OK: output format valid")
        return 1 if problems else 0

    # Архив дополнительных серий (2.4.1, ТЗ п. 2.7). Если каталоги серий не заданы явно, серии пишутся во
    # временный каталог и попадают только в архив; заданные явно каталоги остаются как были и тоже входят в архив.
    series_zip = series_zip_path(output_csv, args.series_zip)
    series_stage: Optional[Path] = None
    sr_study, sr_study_dir, sc_dir, seg_dir = args.sr_study, args.sr_study_dir, args.sc_dir, args.seg_dir
    if series_zip is not None:
        try:
            series_stage = Path(tempfile.mkdtemp(prefix="densito_series_"))
            if not (sr_study or sr_study_dir):
                sr_study, sr_study_dir = True, str(series_stage / "sr")
            sc_dir = sc_dir or str(series_stage / "sc")
            seg_dir = seg_dir or str(series_stage / "seg")
        except Exception as e:  # noqa: BLE001 — архив серий не должен мешать основной выгрузке
            LOG.warning("series zip disabled: %s", e)
            series_zip, series_stage = None, None
    try:
        engine = DensitoInference(cfg=cfg, models_dir=args.models_dir, use_embeddings=not args.no_embeddings,
                                   visualize_dir=args.visualize_dir, sr_dir=args.sr_dir,
                                   roi_autocorrect_dir=args.roi_autocorrect_dir,
                                   sr_study=sr_study, sr_study_dir=sr_study_dir, extras=args.extras,
                                   sr_per_image=args.sr_per_image, sc_dir=sc_dir,
                                   seg_dir=seg_dir)
        debug_csv = None
        if args.debug_csv:
            debug_csv = (output_csv.with_name(output_csv.stem + "_debug.csv") if args.debug_csv == "auto"
                         else Path(args.debug_csv))
        rows = engine.run(Path(args.input), output_csv, debug_csv=debug_csv, xlsx=args.xlsx, limit=args.limit)
        if series_zip is not None:
            res_zip = write_batch_series_zip(engine, rows, series_zip)
            if debug_csv is not None and series_stage is not None:
                relink_debug_series_paths(Path(debug_csv), series_stage, series_zip, (res_zip or {}).get("arc_by_src", {}))
        problems = validate_output_csv(output_csv, cfg)
        if problems:
            LOG.error("Output format problems: %s", problems[:10])
        else:
            LOG.info("Output format check: OK")
        if not rows:
            # Пустая папка, пустой zip или только не-DICOM файлы: CSV с одним заголовком уже записан,
            # но пустой результат — не успех пакета (код 2, как у фатальной ошибки).
            msg = (f"Во входе «{args.input}» не найдено ни одного DICOM-файла (папка или архив пусты либо "
                   f"содержат только не-DICOM файлы). Записан CSV только с заголовком: {output_csv}")
            LOG.error(msg)
            print(msg, file=sys.stderr)
            return 2
        return 0
    except Exception as e:  # noqa: BLE001  — последний рубеж: файл с заголовком всё равно должен быть
        LOG.critical("Fatal error in batch: %s\n%s", e, traceback.format_exc())
        try:
            if not output_csv.exists():
                write_results([], output_csv, cfg)
        except Exception:  # noqa: BLE001
            pass
        return 2
    finally:
        if series_stage is not None:
            shutil.rmtree(series_stage, ignore_errors=True)


def series_zip_path(output_csv: Path, explicit: Optional[str] = None) -> Optional[Path]:
    """Путь архива дополнительных серий: явный --series-zip, иначе <каталог CSV>/additional_series.zip;
    None — архив выключен (DENSITO_SERIES_ZIP=0 и флаг не задан)."""
    if explicit:
        return Path(explicit)
    if os.environ.get("DENSITO_SERIES_ZIP", "1").strip().lower() in ("0", "false", "no", "off"):
        return None
    from series_zip import SERIES_ZIP_NAME
    return Path(output_csv).with_suffix(".csv").parent / SERIES_ZIP_NAME


def relink_debug_series_paths(debug_csv: Path, stage: Path, zip_path: Path, arc_by_src: Dict[str, str]) -> None:
    """Служебный results_debug.csv: пути во временный каталог серий (он удаляется после упаковки) заменить на путь
    внутри additional_series.zip, а для файлов, которых в архиве нет (PNG/JSON сегментации), — очистить. Официальный
    results.csv не затрагивается; ошибка только логируется."""
    try:
        if not Path(debug_csv).is_file():
            return
        prefix = str(Path(stage))
        with open(debug_csv, newline="", encoding="utf-8") as f:
            rd = csv.reader(f)
            data = list(rd)
        changed = 0
        for row in data[1:]:
            for j, v in enumerate(row):
                if v.startswith(prefix):
                    arc = arc_by_src.get(v, "")
                    row[j] = f"{Path(zip_path).name}/{arc}" if arc else ""
                    changed += 1
        if changed:
            tmp = Path(debug_csv).with_suffix(".csv.tmp")
            with open(tmp, "w", newline="", encoding="utf-8") as f:
                csv.writer(f, lineterminator="\n").writerows(data)
            os.replace(tmp, debug_csv)
    except Exception as e:  # noqa: BLE001
        LOG.warning("debug csv series paths not relinked: %s", e)


def write_batch_series_zip(engine: "DensitoInference", rows: List[Dict[str, Any]], zip_path: Path) -> Optional[Dict[str, Any]]:
    """Упаковать дополнительные серии последнего run() в zip. Никогда не бросает исключение: results.csv уже
    записан, ошибка упаковки только логируется (код возврата пакета не меняется)."""
    try:
        import series_zip as _sz
        dbg = list(getattr(engine, "last_debug_rows", []) or [])
        entries = _sz.plan_entries(rows, _sz.image_files_from_debug(dbg), getattr(engine, "last_study_sr", {}) or {},
                                   [str((d or {}).get("internal_region") or "") for d in dbg])
        res = _sz.write_series_zip(entries, Path(zip_path))
        res["arc_by_src"] = {str(e["src"]): str(e["arc"]) for e in entries}
        LOG.info("Additional series zip: %d DICOM (%s) -> %s", res["n_dicom"],
                 ", ".join(f"{k}={v}" for k, v in sorted(res["by_kind"].items())), zip_path)
        return res
    except Exception as e:  # noqa: BLE001
        LOG.warning("additional series zip not written (%s); results.csv is not affected", e)
        return None


if __name__ == "__main__":
    sys.exit(main())
